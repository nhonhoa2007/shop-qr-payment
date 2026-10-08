#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
generate_de_cuong_dacs2.py
Đề cương chi tiết Đồ án cơ sở 2 (ĐACS 2) - Phùng Nhơn Hòa (MSSV: 25ITE020)
Có nêu rõ Tech Stack công nghệ tinh gọn, chuẩn sinh viên CNTT (VKU).
Thuần trắng đen (Black & White Only), vừa vặn chuẩn 2 trang A4 in 2 mặt.
"""

import os
import subprocess
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_margins(cell, top=35, bottom=35, left=70, right=70):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for side, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{side}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="000000", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def remove_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'  <w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'  <w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'  <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'  <w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'  <w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def build_docx(output_docx):
    doc = docx.Document()
    
    # Thiết lập lề A4
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.4)
    sec.bottom_margin = Cm(1.4)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(1.6)

    # Style mặc định thuần đen trắng
    norm = doc.styles['Normal']
    norm.font.name = 'Times New Roman'
    norm.font.size = Pt(10)
    norm.font.color.rgb = RGBColor(0, 0, 0)
    norm.paragraph_format.line_spacing = 1.1
    norm.paragraph_format.space_after = Pt(1.5)
    norm.paragraph_format.space_before = Pt(0)

    # ==================== TRANG 1 ====================
    # Header Trường & Quốc hiệu (Không viền)
    t_hdr = doc.add_table(rows=1, cols=2)
    t_hdr.alignment = WD_TABLE_ALIGNMENT.CENTER
    remove_table_borders(t_hdr)
    t_hdr.autofit = False
    t_hdr.columns[0].width = Cm(8.3)
    t_hdr.columns[1].width = Cm(8.7)

    p0 = t_hdr.cell(0, 0).paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.line_spacing = 1.05
    p0.paragraph_format.space_after = Pt(0)
    r = p0.add_run("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN\nVÀ TRUYỀN THÔNG VIỆT - HÀN\n"); r.font.size = Pt(9.5); r.bold = True
    r = p0.add_run("KHOA KHOA HỌC MÁY TÍNH\n"); r.font.size = Pt(10); r.bold = True
    r = p0.add_run("------------------------"); r.font.size = Pt(8.5); r.bold = False

    p1 = t_hdr.cell(0, 1).paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.line_spacing = 1.05
    p1.paragraph_format.space_after = Pt(0)
    r = p1.add_run("CỘNG HOÀ XÃ HỘI CHỦ NGHĨA VIỆT NAM\n"); r.font.size = Pt(9.5); r.bold = True
    r = p1.add_run("Độc lập - Tự do - Hạnh phúc\n"); r.font.size = Pt(10); r.bold = True
    r = p1.add_run("------------------------------------\n"); r.font.size = Pt(8.5); r.bold = False
    r = p1.add_run("Đà Nẵng, ngày 08 tháng 10 năm 2026"); r.font.size = Pt(9.5); r.italic = True

    # Tiêu đề ĐỀ CƯƠNG ĐỒ ÁN CƠ SỞ 2
    p_t = doc.add_paragraph()
    p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t.paragraph_format.space_before = Pt(5)
    p_t.paragraph_format.space_after = Pt(4)
    r = p_t.add_run("ĐỀ CƯƠNG ĐỒ ÁN CƠ SỞ 2")
    r.font.size = Pt(13)
    r.bold = True

    # 1. Thông tin người thực hiện
    p_s1 = doc.add_paragraph()
    p_s1.paragraph_format.space_before = Pt(2)
    p_s1.paragraph_format.space_after = Pt(1.5)
    r = p_s1.add_run("1. Thông tin người thực hiện:"); r.bold = True; r.font.size = Pt(10.5)

    t_info = doc.add_table(rows=2, cols=2)
    t_info.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_info, color="000000", sz="4")
    t_info.autofit = False
    t_info.columns[0].width = Cm(8.5)
    t_info.columns[1].width = Cm(8.5)

    c00 = t_info.cell(0, 0)
    set_cell_margins(c00, top=35, bottom=35, left=70, right=70)
    p = c00.paragraphs[0]; p.paragraph_format.line_spacing = 1.08; p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Họ và tên sinh viên: "); r.bold = True; p.add_run("Phùng Nhơn Hòa\n")
    r = p.add_run("Ngành: "); r.bold = True; p.add_run("Công nghệ thông tin\n")
    r = p.add_run("Điện thoại: "); r.bold = True; p.add_run("0325266490")

    c01 = t_info.cell(0, 1)
    set_cell_margins(c01, top=35, bottom=35, left=70, right=70)
    p = c01.paragraphs[0]; p.paragraph_format.line_spacing = 1.08; p.paragraph_format.space_after = Pt(0)
    r = p.add_run("MSSV: "); r.bold = True; p.add_run("25ITE020\n")
    r = p.add_run("Lớp: "); r.bold = True; p.add_run("Đồ án cơ sở 2 (ITe)-13\n")
    r = p.add_run("Email: "); r.bold = True; p.add_run("hoapn.25ite@vku.udn.vn")

    c10 = t_info.cell(1, 0)
    set_cell_margins(c10, top=35, bottom=35, left=70, right=70)
    p = c10.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Giảng viên hướng dẫn: "); r.bold = True; p.add_run("TS. Nguyễn Đức Hiển")

    c11 = t_info.cell(1, 1)
    set_cell_margins(c11, top=35, bottom=35, left=70, right=70)
    p = c11.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Hình thức thực hiện: "); r.bold = True; p.add_run("Cá nhân (Không có teammate)")

    # 2. Tên đồ án
    p_s2 = doc.add_paragraph()
    p_s2.paragraph_format.space_before = Pt(3)
    p_s2.paragraph_format.space_after = Pt(2)
    r = p_s2.add_run("2. Tên đồ án: "); r.bold = True; r.font.size = Pt(10.5)
    r = p_s2.add_run("HỆ THỐNG WEBSITE THƯƠNG MẠI ĐIỆN TỬ")
    r.bold = True; r.font.size = Pt(10.5)

    # 3. Mô tả
    p_s3 = doc.add_paragraph()
    p_s3.paragraph_format.space_before = Pt(3)
    p_s3.paragraph_format.space_after = Pt(1)
    r = p_s3.add_run("3. Mô tả: "); r.bold = True; r.font.size = Pt(10.5)
    r = p_s3.add_run("(mô tả tổng quan đề tài, mục tiêu dự kiến sẽ đạt được, kết quả đề tài là gì)")
    r.italic = True; r.font.size = Pt(9)

    p_d1 = doc.add_paragraph()
    p_d1.paragraph_format.space_after = Pt(1)
    r = p_d1.add_run("- Mô tả tổng quan đề tài: "); r.bold = True
    p_d1.add_run(
        "Xây dựng hệ thống website thương mại điện tử phục vụ hoạt động bán hàng trực tuyến; "
        "hỗ trợ khách hàng tìm kiếm sản phẩm, đặt hàng và thanh toán trực tuyến qua mã QR, "
        "đồng thời cung cấp phân hệ quản trị cho nhân viên và quản trị viên quản lý sản phẩm, đơn hàng và khách hàng."
    )

    # Tech stack
    p_tech = doc.add_paragraph()
    p_tech.paragraph_format.space_after = Pt(0.5)
    r = p_tech.add_run("- Tech Stack sử dụng trong đề tài:"); r.bold = True

    tech_bullets = [
        ("Frontend: ", "Next.js, React, Tailwind CSS, TypeScript."),
        ("Backend: ", "Node.js (Next.js API Routes), Prisma ORM."),
        ("Cơ sở dữ liệu: ", "PostgreSQL."),
        ("Tích hợp dịch vụ: ", "Thanh toán mã QR (VietQR), API Giao Hàng Nhanh (GHN)."),
    ]
    for tag, val in tech_bullets:
        p_b = doc.add_paragraph()
        p_b.paragraph_format.left_indent = Cm(0.3)
        p_b.paragraph_format.space_after = Pt(0.5)
        r = p_b.add_run("● "); r.font.size = Pt(8)
        r = p_b.add_run(tag); r.bold = True
        p_b.add_run(val)

    # Mục tiêu dự kiến
    p_obj = doc.add_paragraph()
    p_obj.paragraph_format.space_before = Pt(1)
    p_obj.paragraph_format.space_after = Pt(0.5)
    r = p_obj.add_run("- Mục tiêu dự kiến:"); r.bold = True

    obj_items = [
        "Thiết kế giao diện website hiện đại, thân thiện và hiển thị tốt trên cả máy tính lẫn điện thoại di động (Responsive).",
        "Xây dựng cơ sở dữ liệu quan hệ hoàn chỉnh để lưu trữ dữ liệu sản phẩm, người dùng, giỏ hàng, đơn hàng và phản hồi.",
        "Xây dựng các chức năng cốt lõi: Đăng ký/đăng nhập, tìm kiếm & lọc sản phẩm, giỏ hàng, đặt hàng và thanh toán trực tuyến.",
        "Phân quyền người dùng rõ ràng theo vai trò: Khách hàng mua sắm, Nhân viên (Staff) xử lý đơn/kho, và Quản trị viên (Admin) quản lý toàn diện.",
    ]
    for item in obj_items:
        p_item = doc.add_paragraph()
        p_item.paragraph_format.left_indent = Cm(0.3)
        p_item.paragraph_format.space_after = Pt(0.5)
        r = p_item.add_run("● "); r.font.size = Pt(8)
        p_item.add_run(item)

    # Kết quả đạt được
    p_res = doc.add_paragraph()
    p_res.paragraph_format.space_before = Pt(1)
    p_res.paragraph_format.space_after = Pt(0.5)
    r = p_res.add_run("- Kết quả đạt được:"); r.bold = True

    res_items = [
        "Website thương mại điện tử hoạt động thực tế trên trình duyệt web với đầy đủ các tính năng cho Khách hàng, Nhân viên và Admin.",
        "Báo cáo chi tiết về đồ án cơ sở 2 theo đúng quy định của nhà trường.",
        "Bộ slide trình chiếu phục vụ báo cáo và bảo vệ đồ án.",
    ]
    for item in res_items:
        p_item = doc.add_paragraph()
        p_item.paragraph_format.left_indent = Cm(0.3)
        p_item.paragraph_format.space_after = Pt(0.5)
        r = p_item.add_run("● "); r.font.size = Pt(8)
        p_item.add_run(item)

    # 4. Nội dung thực hiện
    p_s4 = doc.add_paragraph()
    p_s4.paragraph_format.space_before = Pt(2.5)
    p_s4.paragraph_format.space_after = Pt(1)
    r = p_s4.add_run("4. Nội dung thực hiện: "); r.bold = True; r.font.size = Pt(10.5)
    r = p_s4.add_run("(các nội dung sẽ thực hiện trong đề tài)")
    r.italic = True; r.font.size = Pt(9)

    nd_items = [
        ("Phân tích hệ thống: ", "Xác định các tác nhân (Khách hàng, Nhân viên, Admin) và các đối tượng cần quản lý: Người dùng, Sản phẩm, Biến thể sản phẩm (màu sắc/kích thước), Giỏ hàng, Đơn hàng, Đánh giá."),
        ("Thiết kế Cơ sở dữ liệu: ", "Xây dựng sơ đồ thực thể liên kết (ERD), thiết kế các bảng dữ liệu, khóa chính, khóa ngoại và cài đặt CSDL quan hệ PostgreSQL."),
        ("Xây dựng Backend: ", "Viết các hàm và API xử lý nghiệp vụ: Đăng nhập/đăng ký, quản lý sản phẩm, xử lý giỏ hàng, đặt hàng, cập nhật trạng thái đơn hàng và kiểm tra quyền hạn (Role: Customer, Staff, Admin)."),
        ("Lập trình Giao diện: ", "Thiết kế các màn hình tương tác: Trang chủ, Danh mục sản phẩm, Chi tiết sản phẩm, Giỏ hàng, Thanh toán đơn hàng và Dashboard quản trị cho Admin/Nhân viên."),
        ("Kiểm thử & Hoàn thiện: ", "Chạy thử nghiệm các trường hợp kiểm thử (Test cases) để đảm bảo hệ thống chạy ổn định, không lỗi khi nhập sai dữ liệu; hoàn thiện báo cáo và slide thuyết trình."),
    ]
    for tag, body in nd_items:
        p_item = doc.add_paragraph()
        p_item.paragraph_format.left_indent = Cm(0.3)
        p_item.paragraph_format.space_after = Pt(0.5)
        r = p_item.add_run("● "); r.font.size = Pt(8)
        r = p_item.add_run(tag); r.bold = True
        p_item.add_run(body)

    # ==================== NGẮT TRANG CHỦ ĐỘNG ====================
    doc.add_page_break()

    # ==================== TRANG 2: KẾ HOẠCH & CHỮ KÝ ====================
    p_s5 = doc.add_paragraph()
    p_s5.paragraph_format.space_before = Pt(0)
    p_s5.paragraph_format.space_after = Pt(3)
    r = p_s5.add_run("5. Kế hoạch thực hiện: "); r.bold = True; r.font.size = Pt(10.5)
    r = p_s5.add_run("(Dự kiến tiến độ 5 tuần - Sinh viên thực hiện cá nhân)")
    r.italic = True; r.font.size = Pt(9)

    # Bảng kế hoạch 5 tuần chuẩn mẫu ĐACS
    tbl_p = doc.add_table(rows=6, cols=3)
    tbl_p.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_p, color="000000", sz="4")
    tbl_p.autofit = False
    tbl_p.columns[0].width = Cm(3.3)
    tbl_p.columns[1].width = Cm(11.0)
    tbl_p.columns[2].width = Cm(2.9)

    # Header Row (Nền trắng, chữ đen in đậm)
    h_titles = ["Thời gian", "Nội dung thực hiện", "Người phụ trách chính"]
    for i, t in enumerate(h_titles):
        c = tbl_p.rows[0].cells[i]
        set_cell_margins(c, top=45, bottom=45, left=65, right=65)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(t)
        r.bold = True
        r.font.size = Pt(9.5)

    plan_rows = [
        (
            "Tuần 1\n(12/10/2026 -\n18/10/2026)",
            "● Phân tích các thực thể: Khách hàng, Nhân viên, Sản phẩm, Biến thể, Đơn hàng, Giỏ hàng.\n"
            "● Thiết kế sơ đồ quan hệ ERD và cài đặt cơ sở dữ liệu PostgreSQL qua Prisma ORM.\n"
            "● Khởi tạo dự án Next.js, thiết lập cấu trúc mã nguồn và kết nối cơ sở dữ liệu.",
            "Phùng Nhơn Hòa"
        ),
        (
            "Tuần 2\n(19/10/2026 -\n25/10/2026)",
            "● Xây dựng chức năng đăng nhập, đăng ký và phân quyền (Customer, Staff, Admin).\n"
            "● Viết các API/hàm thao tác Thêm, Sửa, Xóa, Lấy danh sách Sản phẩm và Danh mục.\n"
            "● Thiết kế giao diện Trang chủ, Trang danh mục và Bộ lọc tìm kiếm sản phẩm.",
            "Phùng Nhơn Hòa"
        ),
        (
            "Tuần 3\n(26/10/2026 -\n01/11/2026)",
            "● Thiết kế giao diện Trang chi tiết sản phẩm và chức năng Giỏ hàng.\n"
            "● Xây dựng logic nghiệp vụ: Thêm vào giỏ, cập nhật số lượng, kiểm tra tồn kho.\n"
            "● Xây dựng quy trình Tạo đơn hàng và cập nhật trạng thái đơn.",
            "Phùng Nhơn Hòa"
        ),
        (
            "Tuần 4\n(02/11/2026 -\n08/11/2026)",
            "● Tích hợp chức năng hiển thị mã QR thanh toán và xác nhận đơn hàng.\n"
            "● Xây dựng màn hình Dashboard quản trị cho Admin và Nhân viên (Staff).\n"
            "● Viết các truy vấn thống kê doanh thu và báo cáo đơn hàng dạng biểu đồ/bảng.",
            "Phùng Nhơn Hòa"
        ),
        (
            "Tuần 5\n(09/11/2026 -\n15/11/2026)",
            "● Kiểm tra các trường hợp lỗi (nhập sai dữ liệu, hết hàng, quyền truy cập).\n"
            "● Căn chỉnh giao diện và hoàn thiện trải nghiệm người dùng.\n"
            "● Hoàn thành báo cáo chi tiết đồ án và làm bài thuyết trình đồ án.",
            "Phùng Nhơn Hòa"
        ),
    ]

    for idx, (t_col, c_col, a_col) in enumerate(plan_rows):
        cells = tbl_p.rows[idx + 1].cells
        for c in cells:
            set_cell_margins(c, top=35, bottom=35, left=65, right=65)

        # Cột 1: Thời gian
        p = cells[0].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.line_spacing = 1.08; p.paragraph_format.space_after = Pt(0)
        r = p.add_run(t_col); r.font.size = Pt(8.5); r.bold = True

        # Cột 2: Nội dung
        p = cells[1].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.line_spacing = 1.08; p.paragraph_format.space_after = Pt(0)
        r = p.add_run(c_col); r.font.size = Pt(9)

        # Cột 3: Người phụ trách
        p = cells[2].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.line_spacing = 1.08; p.paragraph_format.space_after = Pt(0)
        r = p.add_run(a_col); r.font.size = Pt(9); r.bold = True

    # Chữ ký phê duyệt
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(8)
    p_sp.paragraph_format.space_after = Pt(0)

    tbl_s = doc.add_table(rows=1, cols=2)
    tbl_s.alignment = WD_TABLE_ALIGNMENT.CENTER
    remove_table_borders(tbl_s)
    tbl_s.autofit = False
    tbl_s.columns[0].width = Cm(8.6)
    tbl_s.columns[1].width = Cm(8.6)

    # Giảng viên
    p = tbl_s.cell(0, 0).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.line_spacing = 1.12; p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Ngày ..... tháng ..... năm 2026\n"); r.italic = True; r.font.size = Pt(9.5)
    r = p.add_run("GIẢNG VIÊN HƯỚNG DẪN\n"); r.bold = True; r.font.size = Pt(10)
    r = p.add_run("(ký và ghi rõ họ tên)\n\n\n\n\n"); r.italic = True; r.font.size = Pt(9)
    r = p.add_run("TS. Nguyễn Đức Hiển"); r.bold = True; r.font.size = Pt(10)

    # Sinh viên
    p = tbl_s.cell(0, 1).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.line_spacing = 1.12; p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Đà Nẵng, ngày 08 tháng 10 năm 2026\n"); r.italic = True; r.font.size = Pt(9.5)
    r = p.add_run("SINH VIÊN THỰC HIỆN\n"); r.bold = True; r.font.size = Pt(10)
    r = p.add_run("(ký và ghi rõ họ tên)\n\n\n\n\n"); r.italic = True; r.font.size = Pt(9)
    r = p.add_run("Phùng Nhơn Hòa"); r.bold = True; r.font.size = Pt(10)

    doc.save(output_docx)
    print(f"[OK] Đã xuất file DOCX có Tech Stack: {output_docx}")

if __name__ == "__main__":
    out_path = "/home/nhonhoa/shop-qr-payment/de_cuong_do_an_co_so_2_Phung_Nhon_Hoa.docx"
    build_docx(out_path)
