#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
generate_report_dacs2.py
Xuất bản Báo cáo Đồ án cơ sở 2 (DACS2) chuẩn mẫu VKU DOCX và PDF
Dự án: Shop QR Payment (shop-qr-payment)
Tác giả: @document-specialist (Hermes Agent)
"""

import os
import re
import sys
import shutil
import subprocess
from PIL import Image

import docx
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

# --- CẤU HÌNH ĐƯỜNG DẪN DỰ ÁN ---
PROJECT_ROOT = "/home/nhonhoa/shop-qr-payment"
ACADEMIC_DIR = os.path.join(PROJECT_ROOT, "docs", "academic")
UML_IMAGES_DIR = os.path.join(PROJECT_ROOT, "docs", "uml", "images")
MOCKUPS_DIR = os.path.join(PROJECT_ROOT, "public", "admin-mockups")
BAN_THAO_PATH = os.path.join(ACADEMIC_DIR, "BAN_THAO_BAO_CAO_DACS2.md")
DOCX_OUT_PATH = os.path.join(ACADEMIC_DIR, "BAO_CAO_DACS2_SHOP_QR_PAYMENT.docx")
PDF_OUT_PATH = os.path.join(ACADEMIC_DIR, "BAO_CAO_DACS2_SHOP_QR_PAYMENT.pdf")

# Màu sắc tiêu chuẩn
COLOR_NAVY = RGBColor(30, 58, 138)       # #1E3A8A - Heading 1 & 2, Table Header
COLOR_SLATE_DARK = RGBColor(15, 23, 42)  # #0F172A - Heading 3
COLOR_SLATE_MED = RGBColor(51, 65, 85)   # #334155 - Heading 4
COLOR_BODY = RGBColor(0, 0, 0)           # #000000 - Body Text
COLOR_GRAY = RGBColor(100, 116, 139)     # #64748B - Header & Muted notes
COLOR_CODE = RGBColor(30, 41, 59)        # #1E293B - Code text
HEX_NAVY = "1E3A8A"
HEX_BORDER = "CBD5E1"
HEX_ZEBRA = "F8FAFC"
HEX_WHITE = "FFFFFF"


def set_cell_margins(cell, top=120, bottom=120, left=140, right=140):
    """Thiết lập padding cho ô bảng biểu (đơn vị dxa, 20 dxa = 1 pt)"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for side, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        m = OxmlElement(f'w:{side}')
        m.set(qn('w:w'), str(val))
        m.set(qn('w:type'), 'dxa')
        tcMar.append(m)
    tcPr.append(tcMar)


def set_cell_shading(cell, color_hex):
    """Thiết lập màu nền cho ô"""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)


def set_table_borders(table, color_hex=HEX_BORDER):
    """Thiết lập viền bảng thanh mảnh tiêu chuẩn"""
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for b in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{b}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')  # 0.5 pt
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), color_hex)
        tblBorders.append(border)
    tblPr.append(tblBorders)


def parse_inline_formatting(paragraph, text, base_font_size=13, is_heading=False, base_color=COLOR_BODY, base_italic=False, base_bold=False):
    """Phân tích cú pháp inline markdown (**bold**, *italic*, `code`, <br>) và tạo runs"""
    # Xử lý <br> thành ký tự ngắt dòng đặc biệt
    text = re.sub(r'<br\s*/?>', '\n', text)
    
    # Tokenizer bằng regex
    tokens = re.split(r'(\*\*\*.*?\*\*\*|\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for tok in tokens:
        if not tok:
            continue
        
        bold = base_bold
        italic = base_italic
        code = False
        content = tok
        
        if tok.startswith('***') and tok.endswith('***') and len(tok) >= 6:
            content = tok[3:-3]
            bold = True
            italic = True
        elif tok.startswith('**') and tok.endswith('**') and len(tok) >= 4:
            content = tok[2:-2]
            bold = True
        elif tok.startswith('*') and tok.endswith('*') and len(tok) >= 2:
            content = tok[1:-1]
            italic = True
        elif tok.startswith('`') and tok.endswith('`') and len(tok) >= 2:
            content = tok[1:-1]
            code = True
        
        # Nếu có \n trong content (từ <br> hoặc xuống dòng)
        lines = content.split('\n')
        for i, line in enumerate(lines):
            if i > 0:
                paragraph.add_run().add_break()
            if not line:
                continue
            run = paragraph.add_run(line)
            run.bold = bold
            run.italic = italic
            if code:
                run.font.name = 'Consolas'
                run.font.size = Pt(base_font_size - 1.5)
                run.font.color.rgb = COLOR_CODE
            else:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(base_font_size)
                run.font.color.rgb = base_color


def add_image_scaled(doc, img_path, caption=None, max_w_cm=15.5, max_h_cm=17.5):
    """Chèn hình ảnh với tỷ lệ khung hình tự động co giãn không tràn trang"""
    if not os.path.exists(img_path):
        print(f"  [WARN] Image not found: {img_path}")
        return
    
    try:
        with Image.open(img_path) as im:
            w_px, h_px = im.size
            aspect = w_px / h_px
        
        # Tính kích thước phù hợp
        if aspect >= (max_w_cm / max_h_cm):
            target_w = Cm(max_w_cm)
            target_h = Cm(max_w_cm / aspect)
        else:
            target_h = Cm(max_h_cm)
            target_w = Cm(max_h_cm * aspect)
        
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.first_line_indent = Cm(0)
        
        run = p.add_run()
        run.add_picture(img_path, width=target_w, height=target_h)
        
        if caption:
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(4)
            p_cap.paragraph_format.space_after = Pt(12)
            p_cap.paragraph_format.first_line_indent = Cm(0)
            
            run_cap = p_cap.add_run(caption)
            run_cap.font.name = 'Times New Roman'
            run_cap.font.size = Pt(11)
            run_cap.font.italic = True
            run_cap.font.color.rgb = COLOR_GRAY
    except Exception as e:
        print(f"  [ERR] Failed to insert image {img_path}: {e}")


def build_markdown_table(doc, raw_lines, caption=None):
    """Dựng bảng biểu OpenXML hoàn chỉnh từ khối markdown table"""
    if len(raw_lines) < 2:
        return
    
    # Tách dữ liệu
    header_line = raw_lines[0]
    align_line = raw_lines[1]
    data_lines = raw_lines[2:]
    
    headers = [c.strip() for c in header_line.split('|')[1:-1]]
    num_cols = len(headers)
    if num_cols == 0:
        return
    
    # Xác định căn lề từng cột từ dòng align
    col_aligns = []
    align_tokens = [c.strip() for c in align_line.split('|')[1:-1]]
    for tok in align_tokens:
        if tok.startswith(':') and tok.endswith(':'):
            col_aligns.append(WD_ALIGN_PARAGRAPH.CENTER)
        elif tok.endswith(':'):
            col_aligns.append(WD_ALIGN_PARAGRAPH.RIGHT)
        else:
            col_aligns.append(WD_ALIGN_PARAGRAPH.LEFT)
    while len(col_aligns) < num_cols:
        col_aligns.append(WD_ALIGN_PARAGRAPH.LEFT)
    
    # Tính toán độ rộng cột (Tổng cộng = 16.0 cm)
    TOTAL_WIDTH_CM = 16.0
    if num_cols == 2:
        col_widths_cm = [5.0, 11.0]
    elif num_cols == 4:
        # Kiểm tra nếu là bảng đánh giá GVHD hoặc UC
        if headers[0] == 'STT' and 'Điểm' in headers[2]:
            col_widths_cm = [1.2, 8.8, 3.0, 3.0]
        elif headers[0] == 'Mã UC':
            col_widths_cm = [2.0, 4.0, 3.5, 6.5]
        elif headers[0] == 'STT':
            col_widths_cm = [1.2, 3.5, 4.0, 7.3]
        else:
            col_widths_cm = [TOTAL_WIDTH_CM / num_cols] * num_cols
    elif num_cols == 5:
        # Bảng từ điển CSDL
        col_widths_cm = [1.2, 3.3, 2.7, 3.3, 5.5]
    else:
        col_widths_cm = [TOTAL_WIDTH_CM / num_cols] * num_cols
    
    # Chèn Caption bảng (trên bảng)
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_cap.paragraph_format.space_before = Pt(8)
        p_cap.paragraph_format.space_after = Pt(4)
        p_cap.paragraph_format.first_line_indent = Cm(0)
        
        run_cap = p_cap.add_run(caption)
        run_cap.font.name = 'Times New Roman'
        run_cap.font.size = Pt(11)
        run_cap.font.bold = True
        run_cap.font.color.rgb = COLOR_NAVY
    
    # Tạo bảng
    table = doc.add_table(rows=len(data_lines) + 1, cols=num_cols)
    set_table_borders(table)
    table.autofit = False
    
    # Cấu hình lặp lại Header
    header_tr = table.rows[0]._tr.get_or_add_trPr()
    header_tr.append(OxmlElement('w:tblHeader'))
    header_tr.append(OxmlElement('w:cantSplit'))
    
    # Định dạng Header
    for col_idx, h_text in enumerate(headers):
        cell = table.rows[0].cells[col_idx]
        cell.width = Cm(col_widths_cm[col_idx])
        tcPr = cell._tc.get_or_add_tcPr()
        tcW = OxmlElement('w:tcW')
        tcW.set(qn('w:w'), str(int(col_widths_cm[col_idx] * 567)))
        tcW.set(qn('w:type'), 'dxa')
        tcPr.append(tcW)
        
        set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
        set_cell_shading(cell, HEX_NAVY)
        
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.first_line_indent = Cm(0)
        
        parse_inline_formatting(p, h_text, base_font_size=11, base_color=RGBColor(255, 255, 255), base_bold=True)
    
    # Định dạng Dữ liệu
    for row_idx, row_str in enumerate(data_lines):
        row = table.rows[row_idx + 1]
        row_tr = row._tr.get_or_add_trPr()
        row_tr.append(OxmlElement('w:cantSplit'))
        
        cells_data = [c.strip() for c in row_str.split('|')[1:-1]]
        # Padding nếu thiếu cell
        while len(cells_data) < num_cols:
            cells_data.append('')
        
        bg_color = HEX_ZEBRA if (row_idx % 2 == 1) else HEX_WHITE
        
        for col_idx in range(num_cols):
            cell = row.cells[col_idx]
            cell.width = Cm(col_widths_cm[col_idx])
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = OxmlElement('w:tcW')
            tcW.set(qn('w:w'), str(int(col_widths_cm[col_idx] * 567)))
            tcW.set(qn('w:type'), 'dxa')
            tcPr.append(tcW)
            
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            set_cell_shading(cell, bg_color)
            
            p = cell.paragraphs[0]
            # STT hoặc cột ngắn thì center, còn lại theo col_align
            if col_idx == 0 and num_cols > 2 and len(cells_data[col_idx]) <= 4:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = col_aligns[col_idx]
            
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.first_line_indent = Cm(0)
            
            parse_inline_formatting(p, cells_data[col_idx], base_font_size=10.5, base_color=COLOR_BODY)

    # Thêm khoảng cách sau bảng
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)
    p_after.paragraph_format.first_line_indent = Cm(0)


def build_outer_cover(doc):
    """Trang bìa chính (Outer Cover) chuẩn đại học VKU"""
    p_top = doc.add_paragraph()
    p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_top.paragraph_format.space_before = Pt(12)
    p_top.paragraph_format.space_after = Pt(4)
    p_top.paragraph_format.first_line_indent = Cm(0)
    
    run_sch = p_top.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO — ĐẠI HỌC ĐÀ NẴNG\n")
    run_sch.font.name = "Times New Roman"
    run_sch.font.size = Pt(12.5)
    run_sch.font.bold = True
    
    run_vku = p_top.add_run("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN VÀ TRUYỀN THÔNG VIỆT - HÀN\n")
    run_vku.font.name = "Times New Roman"
    run_vku.font.size = Pt(13)
    run_vku.font.bold = True
    run_vku.font.color.rgb = COLOR_NAVY
    
    run_fac = p_top.add_run("KHOA KHOA HỌC MÁY TÍNH\n")
    run_fac.font.name = "Times New Roman"
    run_fac.font.size = Pt(12)
    run_fac.font.bold = True
    
    run_line = p_top.add_run("______________________________\n")
    run_line.font.name = "Times New Roman"
    run_line.font.size = Pt(11)
    run_line.font.bold = True
    run_line.font.color.rgb = COLOR_GRAY
    
    # Khoảng cách giữa trang
    for _ in range(4):
        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(0)
        p_sp.paragraph_format.space_after = Pt(0)
        p_sp.paragraph_format.first_line_indent = Cm(0)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(14)
    p_title.paragraph_format.first_line_indent = Cm(0)
    
    run_rep = p_title.add_run("BÁO CÁO ĐỒ ÁN CƠ SỞ 2\n\n")
    run_rep.font.name = "Times New Roman"
    run_rep.font.size = Pt(20)
    run_rep.font.bold = True
    run_rep.font.color.rgb = COLOR_NAVY
    
    run_proj = p_title.add_run("ĐỀ TÀI:\nWEBSITE THƯƠNG MẠI ĐIỆN TỬ TÍCH HỢP THANH TOÁN VIETQR TỰ ĐỘNG\nVÀ QUẢN TRỊ SAAS BENTO GRID\n(SHOP QR PAYMENT)")
    run_proj.font.name = "Times New Roman"
    run_proj.font.size = Pt(14.5)
    run_proj.font.bold = True
    run_proj.font.color.rgb = COLOR_NAVY
    
    for _ in range(4):
        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(0)
        p_sp.paragraph_format.space_after = Pt(0)
        p_sp.paragraph_format.first_line_indent = Cm(0)

    # Khung thông tin sinh viên & giảng viên
    p_info = doc.add_paragraph()
    p_info.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_info.paragraph_format.left_indent = Cm(4.0)
    p_info.paragraph_format.space_before = Pt(0)
    p_info.paragraph_format.space_after = Pt(2)
    p_info.paragraph_format.line_spacing = 1.35
    p_info.paragraph_format.first_line_indent = Cm(0)
    
    lines_info = [
        ("Sinh viên thực hiện : ", "[Họ và tên sinh viên 1 - MSSV 1]"),
        ("                      ", "[Họ và tên sinh viên 2 - MSSV 2]"),
        ("Lớp                 : ", "[Lớp sinh hoạt / học phần]"),
        ("Giảng viên hướng dẫn: ", "[Học hàm, Học vị & Họ tên Giảng viên hướng dẫn]")
    ]
    for label, val in lines_info:
        r_lbl = p_info.add_run(label)
        r_lbl.font.name = "Times New Roman"
        r_lbl.font.size = Pt(13)
        r_lbl.font.bold = True
        
        r_val = p_info.add_run(val + "\n")
        r_val.font.name = "Times New Roman"
        r_val.font.size = Pt(13)
    
    for _ in range(5):
        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(0)
        p_sp.paragraph_format.space_after = Pt(0)
        p_sp.paragraph_format.first_line_indent = Cm(0)

    p_bot = doc.add_paragraph()
    p_bot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_bot.paragraph_format.space_before = Pt(0)
    p_bot.paragraph_format.space_after = Pt(0)
    p_bot.paragraph_format.first_line_indent = Cm(0)
    
    r_bot = p_bot.add_run("Đà Nẵng, tháng 10 năm 2026")
    r_bot.font.name = "Times New Roman"
    r_bot.font.size = Pt(12)
    r_bot.font.italic = True
    
    doc.add_page_break()


def build_inner_cover(doc):
    """Trang bìa phụ (Inner Cover)"""
    p_top = doc.add_paragraph()
    p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_top.paragraph_format.space_before = Pt(12)
    p_top.paragraph_format.space_after = Pt(4)
    p_top.paragraph_format.first_line_indent = Cm(0)
    
    run_vku = p_top.add_run("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN VÀ TRUYỀN THÔNG VIỆT - HÀN\n")
    run_vku.font.name = "Times New Roman"
    run_vku.font.size = Pt(13)
    run_vku.font.bold = True
    run_vku.font.color.rgb = COLOR_NAVY
    
    run_fac = p_top.add_run("KHOA KHOA HỌC MÁY TÍNH\n")
    run_fac.font.name = "Times New Roman"
    run_fac.font.size = Pt(12)
    run_fac.font.bold = True
    
    run_line = p_top.add_run("______________________________\n")
    run_line.font.name = "Times New Roman"
    run_line.font.size = Pt(11)
    run_line.font.color.rgb = COLOR_GRAY
    
    for _ in range(4):
        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.first_line_indent = Cm(0)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.first_line_indent = Cm(0)
    
    run_rep = p_title.add_run("BÁO CÁO ĐỒ ÁN CƠ SỞ 2\n\n")
    run_rep.font.name = "Times New Roman"
    run_rep.font.size = Pt(20)
    run_rep.font.bold = True
    run_rep.font.color.rgb = COLOR_NAVY
    
    run_proj = p_title.add_run("ĐỀ TÀI:\nWEBSITE THƯƠNG MẠI ĐIỆN TỬ TÍCH HỢP THANH TOÁN VIETQR TỰ ĐỘNG\nVÀ QUẢN TRỊ SAAS BENTO GRID\n(SHOP QR PAYMENT)")
    run_proj.font.name = "Times New Roman"
    run_proj.font.size = Pt(14.5)
    run_proj.font.bold = True
    run_proj.font.color.rgb = COLOR_NAVY
    
    for _ in range(4):
        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.first_line_indent = Cm(0)

    p_info = doc.add_paragraph()
    p_info.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_info.paragraph_format.left_indent = Cm(4.0)
    p_info.paragraph_format.line_spacing = 1.35
    p_info.paragraph_format.first_line_indent = Cm(0)
    
    lines_info = [
        ("Sinh viên thực hiện : ", "[Họ và tên sinh viên 1 - MSSV 1]"),
        ("                      ", "[Họ và tên sinh viên 2 - MSSV 2]"),
        ("Lớp                 : ", "[Lớp sinh hoạt / học phần]"),
        ("Giảng viên hướng dẫn: ", "[Học hàm, Học vị & Họ tên Giảng viên hướng dẫn]")
    ]
    for label, val in lines_info:
        r_lbl = p_info.add_run(label)
        r_lbl.font.name = "Times New Roman"
        r_lbl.font.size = Pt(13)
        r_lbl.font.bold = True
        
        r_val = p_info.add_run(val + "\n")
        r_val.font.name = "Times New Roman"
        r_val.font.size = Pt(13)
    
    for _ in range(5):
        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.first_line_indent = Cm(0)

    p_bot = doc.add_paragraph()
    p_bot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_bot.paragraph_format.first_line_indent = Cm(0)
    
    r_bot = p_bot.add_run("Đà Nẵng, tháng 10 năm 2026")
    r_bot.font.name = "Times New Roman"
    r_bot.font.size = Pt(12)
    r_bot.font.italic = True
    
    doc.add_page_break()


def build_acknowledgements(doc):
    """Trang Lời cảm ơn"""
    p_h = doc.add_paragraph()
    p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_h.paragraph_format.space_before = Pt(18)
    p_h.paragraph_format.space_after = Pt(18)
    p_h.paragraph_format.first_line_indent = Cm(0)
    
    r_h = p_h.add_run("LỜI CẢM ƠN")
    r_h.font.name = "Times New Roman"
    r_h.font.size = Pt(16)
    r_h.font.bold = True
    r_h.font.color.rgb = COLOR_NAVY
    
    p_note = doc.add_paragraph()
    p_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_note.paragraph_format.space_before = Pt(0)
    p_note.paragraph_format.space_after = Pt(12)
    p_note.paragraph_format.first_line_indent = Cm(0)
    r_note = p_note.add_run("[Phần thông tin cá nhân: Sinh viên tự hoàn thành theo thực tế]")
    r_note.font.name = "Times New Roman"
    r_note.font.size = Pt(11)
    r_note.font.italic = True
    r_note.font.color.rgb = COLOR_GRAY

    paragraphs = [
        "Em xin bày tỏ lòng biết ơn sâu sắc nhất đến Quý Thầy/Cô trường Đại học Công nghệ Thông tin & Truyền thông Việt - Hàn (VKU), đặc biệt là các Thầy/Cô Khoa Khoa học Máy tính đã tận tình giảng dạy, truyền đạt những kiến thức quý báu về nền tảng khoa học máy tính, kỹ thuật phần mềm và kiến trúc hệ thống trong suốt thời gian học tập.",
        "Đặc biệt, em xin gửi lời tri ân chân thành nhất tới [Học hàm, Học vị & Họ tên Giảng viên hướng dẫn], người đã trực tiếp hướng dẫn, định hướng chuyên môn và tận tâm đồng hành cùng em trong suốt quá trình triển khai Đồ án cơ sở 2. Những ý kiến đóng góp, chỉ bảo nghiêm cẩn và tâm huyết của Thầy/Cô là kim chỉ nam quan trọng giúp em hoàn thiện đồ án này một cách khoa học, chuyên nghiệp và ứng dụng thực tiễn cao.",
        "Mặc dù đã nỗ lực hết mình để xây dựng hệ thống hoàn chỉnh và vận hành ổn định, song do giới hạn về mặt thời gian và kinh nghiệm thực tế, báo cáo khó tránh khỏi những thiếu sót nhất định. Em rất mong nhận được những nhận xét, góp ý quý báu từ Quý Thầy/Cô trong hội đồng chấm đồ án để sản phẩm được hoàn thiện hơn nữa.",
    ]
    for p_txt in paragraphs:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.35
        p.paragraph_format.first_line_indent = Cm(1.0)
        
        parse_inline_formatting(p, p_txt, base_font_size=13, base_color=COLOR_BODY)

    p_end = doc.add_paragraph()
    p_end.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_end.paragraph_format.space_before = Pt(12)
    p_end.paragraph_format.space_after = Pt(0)
    p_end.paragraph_format.first_line_indent = Cm(0)
    r_end = p_end.add_run("Em xin chân thành cảm ơn!")
    r_end.font.name = "Times New Roman"
    r_end.font.size = Pt(13)
    r_end.font.italic = True
    r_end.font.bold = True
    
    doc.add_page_break()


def build_supervisor_review(doc):
    """Trang Nhận xét của Giảng viên hướng dẫn kèm Bảng tiêu chí"""
    p_h = doc.add_paragraph()
    p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_h.paragraph_format.space_before = Pt(18)
    p_h.paragraph_format.space_after = Pt(14)
    p_h.paragraph_format.first_line_indent = Cm(0)
    
    r_h = p_h.add_run("NHẬN XÉT CỦA GIẢNG VIÊN HƯỚNG DẪN")
    r_h.font.name = "Times New Roman"
    r_h.font.size = Pt(16)
    r_h.font.bold = True
    r_h.font.color.rgb = COLOR_NAVY

    sections = [
        ("1. Ý thức thực hiện đồ án của sinh viên:", 3),
        ("2. Đánh giá nội dung và kết quả đạt được:", 3),
        ("3. Khả năng ứng dụng và tính sáng tạo:", 3),
    ]
    dot_line = "…………………………………………………………………………………………………………"
    
    for title, num_lines in sections:
        p_sec = doc.add_paragraph()
        p_sec.paragraph_format.space_before = Pt(6)
        p_sec.paragraph_format.space_after = Pt(2)
        p_sec.paragraph_format.first_line_indent = Cm(0)
        r_sec = p_sec.add_run(title)
        r_sec.font.name = "Times New Roman"
        r_sec.font.size = Pt(13)
        r_sec.font.bold = True
        
        for _ in range(num_lines):
            p_dot = doc.add_paragraph()
            p_dot.paragraph_format.space_before = Pt(0)
            p_dot.paragraph_format.space_after = Pt(4)
            p_dot.paragraph_format.first_line_indent = Cm(0)
            r_d = p_dot.add_run(dot_line)
            r_d.font.name = "Times New Roman"
            r_d.font.size = Pt(11)
            r_d.font.color.rgb = COLOR_GRAY

    # 4. Điểm đánh giá chi tiết
    p_t = doc.add_paragraph()
    p_t.paragraph_format.space_before = Pt(8)
    p_t.paragraph_format.space_after = Pt(4)
    p_t.paragraph_format.first_line_indent = Cm(0)
    r_t = p_t.add_run("4. Điểm đánh giá chi tiết:")
    r_t.font.name = "Times New Roman"
    r_t.font.size = Pt(13)
    r_t.font.bold = True

    eval_table_lines = [
        "| STT | Tiêu chí đánh giá | Điểm tối đa | Điểm đánh giá |",
        "| :---: | :--- | :---: | :---: |",
        "| 1 | Ý thức thực hiện, tinh thần học hỏi và thái độ | 2.0 | |",
        "| 2 | Khả năng thu thập, nghiên cứu tài liệu và phân tích yêu cầu | 2.0 | |",
        "| 3 | Mức độ hoàn thành khối lượng công việc và chất lượng phần mềm | 4.0 | |",
        "| 4 | Bố cục, hình thức trình bày quyển báo cáo và thuyết minh | 2.0 | |",
        "| **Tổng** | **Tổng điểm đánh giá** | **10.0** | |"
    ]
    build_markdown_table(doc, eval_table_lines)

    # Chữ ký GVHD
    p_sig = doc.add_paragraph()
    p_sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sig.paragraph_format.space_before = Pt(12)
    p_sig.paragraph_format.space_after = Pt(2)
    p_sig.paragraph_format.first_line_indent = Cm(0)
    r_sig1 = p_sig.add_run("Đà Nẵng, ngày … tháng … năm 2026\n")
    r_sig1.font.name = "Times New Roman"
    r_sig1.font.size = Pt(12)
    r_sig1.font.italic = True
    
    r_sig2 = p_sig.add_run("GIẢNG VIÊN HƯỚNG DẪN\n")
    r_sig2.font.name = "Times New Roman"
    r_sig2.font.size = Pt(13)
    r_sig2.font.bold = True
    
    r_sig3 = p_sig.add_run("(Ký và ghi rõ họ tên)\n\n\n\n")
    r_sig3.font.name = "Times New Roman"
    r_sig3.font.size = Pt(11)
    r_sig3.font.italic = True

    doc.add_page_break()


def build_table_of_contents(doc, toc_entries=None):
    """Trang Mục lục chuẩn dot-leader"""
    p_h = doc.add_paragraph()
    p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_h.paragraph_format.space_before = Pt(18)
    p_h.paragraph_format.space_after = Pt(18)
    p_h.paragraph_format.first_line_indent = Cm(0)
    
    r_h = p_h.add_run("MỤC LỤC")
    r_h.font.name = "Times New Roman"
    r_h.font.size = Pt(16)
    r_h.font.bold = True
    r_h.font.color.rgb = COLOR_NAVY

    # Trường Word Field Native TOC
    p_native = doc.add_paragraph()
    p_native.paragraph_format.space_before = Pt(0)
    p_native.paragraph_format.space_after = Pt(4)
    p_native.paragraph_format.first_line_indent = Cm(0)
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'TOC \\o "1-3" \\h \\z \\u')
    p_native._p.append(fldSimple)

    if not toc_entries:
        # Mục lục mặc định dự phòng
        toc_entries = [
            ("CHƯƠNG I: TỔNG QUAN ĐỀ TÀI", 1, 1),
            ("1.1. Giới thiệu", 2, 1),
            ("1.1.1. Mục đích", 3, 1),
            ("1.1.2. Quy ước tài liệu", 3, 1),
            ("1.1.3. Đối tượng và gợi ý", 3, 2),
            ("1.1.4. Phạm vi dự án", 3, 2),
            ("1.2. Mô tả tổng quan", 2, 2),
            ("1.2.1. Quan điểm về hệ thống", 3, 2),
            ("1.2.2. Các lớp người dùng và đặc điểm", 3, 3),
            ("1.2.3. Môi trường hoạt động", 3, 3),
            ("1.2.4. Hạn chế về thiết kế và triển khai", 3, 3),
            ("1.2.5. Tài liệu hướng dẫn người dùng", 3, 4),
            ("1.2.6. Giả định và phụ thuộc", 3, 4),
            ("1.3. Observation (Biên bản họp nhóm)", 2, 4),
            ("CHƯƠNG II: MÔ HÌNH VÀ CHỨC NĂNG", 1, 5),
            ("2.1. Các yêu cầu chức năng", 2, 5),
            ("2.1.1. Các tác nhân", 3, 5),
            ("2.1.2. Các tính năng của hệ thống", 3, 5),
            ("2.1.3. Thành lập các chức năng", 3, 6),
            ("2.1.4. Đặc tả use case", 3, 7),
            ("2.2. Biểu đồ Use Case", 2, 38),
            ("2.3. Class Diagram", 2, 40),
            ("2.4. Activity Diagram", 2, 42),
            ("2.5. Sequence Diagram", 2, 48),
            ("2.6. Communication Diagram", 2, 56),
            ("2.7. Mô hình cơ sở dữ liệu quan hệ", 2, 57),
            ("CHƯƠNG III: XÂY DỰNG ỨNG DỤNG", 1, 74),
            ("3.1. Kiến trúc hệ thống và công nghệ sử dụng", 2, 74),
            ("3.2. Giao diện trang Khách hàng (Storefront)", 2, 77),
            ("3.3. Giao diện trang Quản trị (Admin Bento Grid)", 2, 85),
            ("3.4. Giao diện trang Nhân viên (Staff Portal)", 2, 91),
            ("CHƯƠNG IV: KẾT LUẬN", 1, 93),
            ("4.1. Ưu điểm của hệ thống", 2, 93),
            ("4.2. Hạn chế của hệ thống", 2, 94),
            ("4.3. Hướng phát triển trong tương lai", 2, 95),
            ("4.4. Lời kết", 2, 96),
            ("CHƯƠNG V: TÀI LIỆU THAM KHẢO", 1, 97),
        ]

    for title, level, page_num in toc_entries:
        p_item = doc.add_paragraph()
        p_item.paragraph_format.tab_stops.add_tab_stop(Cm(16.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        p_item.paragraph_format.space_before = Pt(2)
        p_item.paragraph_format.space_after = Pt(2)
        p_item.paragraph_format.first_line_indent = Cm(0)
        
        if level == 1:
            p_item.paragraph_format.left_indent = Cm(0)
            run = p_item.add_run(f"{title}\t{page_num}")
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = COLOR_NAVY
        elif level == 2:
            p_item.paragraph_format.left_indent = Cm(0.6)
            run = p_item.add_run(f"{title}\t{page_num}")
            run.font.name = "Times New Roman"
            run.font.size = Pt(11.5)
            run.font.bold = True
        elif level == 3:
            p_item.paragraph_format.left_indent = Cm(1.2)
            run = p_item.add_run(f"{title}\t{page_num}")
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)
        else:
            p_item.paragraph_format.left_indent = Cm(1.8)
            run = p_item.add_run(f"{title}\t{page_num}")
            run.font.name = "Times New Roman"
            run.font.size = Pt(10.5)
            run.font.italic = True


def build_report_body(doc):
    """Xây dựng 5 Chương nội dung chính từ BAN_THAO_BAO_CAO_DACS2.md"""
    with open(BAN_THAO_PATH, 'r', encoding='utf-8') as f:
        full_text = f.read()

    # Tìm vị trí bắt đầu Chương 1
    ch1_idx = full_text.find("# CHƯƠNG I: TỔNG QUAN ĐỀ TÀI")
    if ch1_idx == -1:
        raise ValueError("Không tìm thấy '# CHƯƠNG I: TỔNG QUAN ĐỀ TÀI' trong bản thảo!")

    body_text = full_text[ch1_idx:]
    lines = body_text.splitlines()

    i = 0
    total_lines = len(lines)
    
    # Biến đếm thứ tự Bảng và Hình tự động
    table_count = {1: 0, 2: 0, 3: 0, 4: 0, 5: 0}
    fig_count = {1: 0, 2: 0, 3: 0, 4: 0, 5: 0}
    current_chapter = 1

    # Mapping ảnh vào các mục
    inserted_images = set()

    while i < total_lines:
        line = lines[i]
        line_strip = line.strip()

        # Bỏ qua dòng trống
        if not line_strip:
            i += 1
            continue

        # Phân tách Heading
        if line_strip.startswith('# '):
            h_text = line_strip[2:].strip()
            # Cập nhật số chương
            if "CHƯƠNG I:" in h_text: current_chapter = 1
            elif "CHƯƠNG II:" in h_text: current_chapter = 2
            elif "CHƯƠNG III:" in h_text: current_chapter = 3
            elif "CHƯƠNG IV:" in h_text: current_chapter = 4
            elif "CHƯƠNG V:" in h_text: current_chapter = 5

            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.first_line_indent = Cm(0)
            # Không ngắt trang nếu là chương 1 (đã ngắt từ section break)
            if current_chapter > 1:
                p.paragraph_format.page_break_before = True

            run = p.add_run(h_text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = COLOR_NAVY
            i += 1
            continue

        elif line_strip.startswith('## '):
            h_text = line_strip[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.first_line_indent = Cm(0)

            run = p.add_run(h_text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = COLOR_NAVY
            i += 1
            continue

        elif line_strip.startswith('### '):
            h_text = line_strip[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.first_line_indent = Cm(0)

            run = p.add_run(h_text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = COLOR_SLATE_DARK
            i += 1

            # Chèn hình ảnh đặc thù nếu gặp mục tương ứng
            if "2.2.1. Biểu đồ Use Case tổng quát" in h_text and "01-use-case" not in inserted_images:
                fig_count[2] += 1
                add_image_scaled(doc, os.path.join(UML_IMAGES_DIR, "01-use-case.png"),
                                 f"Hình 2.{fig_count[2]}: Biểu đồ Use Case tổng quát hệ thống Shop QR Payment (UML 01)")
                inserted_images.add("01-use-case")

            elif "2.3. Class Diagram" in h_text and "02-class-domain" not in inserted_images:
                fig_count[2] += 1
                add_image_scaled(doc, os.path.join(UML_IMAGES_DIR, "02-class-domain.png"),
                                 f"Hình 2.{fig_count[2]}: Sơ đồ lớp thực thể nghiệp vụ và dịch vụ miền (Domain Class Model - UML 02)")
                inserted_images.add("02-class-domain")

            elif "3.1.1. Kiến trúc phân tầng sạch" in h_text and "11-arch" not in inserted_images:
                fig_count[3] += 1
                add_image_scaled(doc, os.path.join(UML_IMAGES_DIR, "11-component-clean-architecture.png"),
                                 f"Hình 3.{fig_count[3]}: Kiến trúc phân tầng sạch Clean Architecture 3 lớp (UML 11)")
                fig_count[3] += 1
                add_image_scaled(doc, os.path.join(UML_IMAGES_DIR, "13-package-modularity.png"),
                                 f"Hình 3.{fig_count[3]}: Sơ đồ gói và cấu trúc module phần mềm (Package Diagram - UML 13)")
                inserted_images.add("11-arch")

            elif "3.1.2. Công nghệ và thư viện cốt lõi" in h_text and "12-deploy" not in inserted_images:
                fig_count[3] += 1
                add_image_scaled(doc, os.path.join(UML_IMAGES_DIR, "12-deployment-infrastructure.png"),
                                 f"Hình 3.{fig_count[3]}: Sơ đồ triển khai hạ tầng đám mây và cơ sở dữ liệu (Deployment Architecture - UML 12)")
                inserted_images.add("12-deploy")

            elif "3.3. Giao diện trang Quản trị" in h_text and "admin-mockups" not in inserted_images:
                fig_count[3] += 1
                add_image_scaled(doc, os.path.join(MOCKUPS_DIR, "admin-options-overview.png"),
                                 f"Hình 3.{fig_count[3]}: Tổng quan các phương án thiết kế Giao diện Quản trị Bento Grid")
                fig_count[3] += 1
                add_image_scaled(doc, os.path.join(MOCKUPS_DIR, "option-1-saas-bento.png"),
                                 f"Hình 3.{fig_count[3]}: Giao diện Bento Grid - 21 chỉ số Analytics thời gian thực")
                inserted_images.add("admin-mockups")

            continue

        elif line_strip.startswith('#### '):
            h_text = line_strip[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.first_line_indent = Cm(0)

            run = p.add_run(h_text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.italic = True
            run.font.color.rgb = COLOR_SLATE_MED
            i += 1
            continue

        elif line_strip.startswith('---'):
            i += 1
            continue

        # Xử lý Bảng biểu (Markdown Table)
        elif line_strip.startswith('|'):
            tbl_lines = []
            while i < total_lines and lines[i].strip().startswith('|'):
                tbl_lines.append(lines[i].strip())
                i += 1
            
            # Đặt Caption bảng tự động
            table_count[current_chapter] += 1
            cap_num = f"Bảng {current_chapter}.{table_count[current_chapter]}"
            
            # Đặt tên bảng dựa trên ngữ cảnh
            first_row = tbl_lines[0]
            if "Thuộc tính" in first_row:
                caption = f"{cap_num}: Bảng thuộc tính đặc tả Use Case"
            elif "Action of Actor" in first_row:
                caption = f"{cap_num}: Kịch bản các bước thực thi Use Case"
            elif "Mã UC" in first_row:
                caption = f"{cap_num}: Danh sách 22 chức năng hệ thống Shop QR Payment"
            elif "Tên Model" in first_row:
                caption = f"{cap_num}: Danh sách 19 bảng dữ liệu trong hệ thống CSDL"
            elif "Tên thuộc tính" in first_row:
                caption = f"{cap_num}: Từ điển dữ liệu thuộc tính và ràng buộc"
            else:
                caption = f"{cap_num}: Bảng dữ liệu chi tiết"

            build_markdown_table(doc, tbl_lines, caption=caption)
            continue

        # Xử lý Danh sách không thứ tự (Bullet points)
        elif line_strip.startswith('- ') or line_strip.startswith('* '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.35
            p.paragraph_format.first_line_indent = Cm(-0.4)

            # Bullet run
            r_b = p.add_run("•  ")
            r_b.font.name = "Times New Roman"
            r_b.font.size = Pt(13)
            r_b.font.bold = True
            r_b.font.color.rgb = COLOR_NAVY

            parse_inline_formatting(p, line_strip[2:].strip(), base_font_size=13, base_color=COLOR_BODY)
            i += 1
            continue

        # Xử lý Danh sách có thứ tự (Numbered list: 1., 2., ...)
        elif re.match(r'^\d+\.\s+', line_strip):
            m = re.match(r'^(\d+\.)\s+(.*)$', line_strip)
            num_prefix = m.group(1)
            item_content = m.group(2)

            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.35
            p.paragraph_format.first_line_indent = Cm(-0.5)

            r_n = p.add_run(num_prefix + "  ")
            r_n.font.name = "Times New Roman"
            r_n.font.size = Pt(13)
            r_n.font.bold = True
            r_n.font.color.rgb = COLOR_NAVY

            parse_inline_formatting(p, item_content, base_font_size=13, base_color=COLOR_BODY)
            i += 1

            # Chèn sơ đồ hoạt động nếu đang ở mục 2.4
            if current_chapter == 2 and "Quy trình Đặt hàng & Giữ kho nguyên tử" in item_content and "07-act" not in inserted_images:
                fig_count[2] += 1
                add_image_scaled(doc, os.path.join(UML_IMAGES_DIR, "07-activity-order-placement.png"),
                                 f"Hình 2.{fig_count[2]}: Sơ đồ hoạt động - Quy trình Đặt hàng & Giữ kho nguyên tử (UML 07)")
                fig_count[2] += 1
                add_image_scaled(doc, os.path.join(UML_IMAGES_DIR, "09-state-machine-orders.png"),
                                 f"Hình 2.{fig_count[2]}: Sơ đồ máy trạng thái - Vòng đời Đơn hàng & Quản lý Kho hàng (UML 09)")
                inserted_images.add("07-act")

            elif current_chapter == 2 and "Bộ lọc An ninh & Đối soát Webhook" in item_content and "08-act" not in inserted_images:
                fig_count[2] += 1
                add_image_scaled(doc, os.path.join(UML_IMAGES_DIR, "08-activity-webhook-security.png"),
                                 f"Hình 2.{fig_count[2]}: Sơ đồ hoạt động - Bộ lọc An ninh & Đối soát Webhook Ngân hàng (UML 08)")
                fig_count[2] += 1
                add_image_scaled(doc, os.path.join(UML_IMAGES_DIR, "15-flowchart-webhook-decision.png"),
                                 f"Hình 2.{fig_count[2]}: Sơ đồ hoạt động - Cây quyết định khớp tiền Webhook (UML 15)")
                inserted_images.add("08-act")

            elif current_chapter == 2 and "Thanh toán Ví Shop 1-chạm CAS" in item_content and "17-act" not in inserted_images:
                fig_count[2] += 1
                add_image_scaled(doc, os.path.join(UML_IMAGES_DIR, "17-activity-wallet-pay.png"),
                                 f"Hình 2.{fig_count[2]}: Sơ đồ hoạt động - Thanh toán Ví Shop 1-chạm CAS nguyên tử (UML 17)")
                inserted_images.add("17-act")

            elif current_chapter == 2 and "Chuỗi dự phòng Định vị GPS" in item_content and "18-act" not in inserted_images:
                fig_count[2] += 1
                add_image_scaled(doc, os.path.join(UML_IMAGES_DIR, "18-flowchart-geolocation.png"),
                                 f"Hình 2.{fig_count[2]}: Sơ đồ hoạt động - Chuỗi dự phòng Định vị GPS sang GHN (UML 18)")
                inserted_images.add("18-act")

            # Chèn sơ đồ tuần tự nếu đang ở mục 2.5
            elif current_chapter == 2 and "UML 03: Đặt hàng & Thanh toán VietQR" in item_content and "03-seq" not in inserted_images:
                fig_count[2] += 1
                add_image_scaled(doc, os.path.join(UML_IMAGES_DIR, "03-sequence-vietqr-payment.png"),
                                 f"Hình 2.{fig_count[2]}: Sơ đồ tuần tự - Đặt hàng & Thanh toán qua mã VietQR động (UML 03)")
                inserted_images.add("03-seq")

            elif current_chapter == 2 and "UML 04: Nạp tiền Ví nội bộ" in item_content and "04-seq" not in inserted_images:
                fig_count[2] += 1
                add_image_scaled(doc, os.path.join(UML_IMAGES_DIR, "04-sequence-wallet-topup.png"),
                                 f"Hình 2.{fig_count[2]}: Sơ đồ tuần tự - Nạp tiền vào Ví Shop qua VietQR tĩnh (UML 04)")
                inserted_images.add("04-seq")

            elif current_chapter == 2 and "UML 05: Hủy đơn & Hoàn tiền Ví CAS" in item_content and "05-seq" not in inserted_images:
                fig_count[2] += 1
                add_image_scaled(doc, os.path.join(UML_IMAGES_DIR, "05-sequence-wallet-refund.png"),
                                 f"Hình 2.{fig_count[2]}: Sơ đồ tuần tự - Hủy đơn & Hoàn tiền Ví CAS nguyên tử (UML 05)")
                fig_count[2] += 1
                add_image_scaled(doc, os.path.join(UML_IMAGES_DIR, "16-sequence-cancel-refund.png"),
                                 f"Hình 2.{fig_count[2]}: Sơ đồ tuần tự - Hủy đơn: Hoàn kho, Thu hồi coupon & Hoàn tiền CAS (UML 16)")
                inserted_images.add("05-seq")

            elif current_chapter == 2 and "UML 06: Chat tư vấn trực tuyến" in item_content and "06-seq" not in inserted_images:
                fig_count[2] += 1
                add_image_scaled(doc, os.path.join(UML_IMAGES_DIR, "06-sequence-realtime-chat.png"),
                                 f"Hình 2.{fig_count[2]}: Sơ đồ tuần tự - Chat tư vấn trực tuyến Real-time theo đơn hàng (UML 06)")
                inserted_images.add("06-seq")

            elif current_chapter == 2 and "UML 14: Vòng đời đầy đủ Đơn hàng" in item_content and "14-seq" not in inserted_images:
                fig_count[2] += 1
                add_image_scaled(doc, os.path.join(UML_IMAGES_DIR, "14-sequence-order-e2e.png"),
                                 f"Hình 2.{fig_count[2]}: Sơ đồ tuần tự - Vòng đời đầy đủ Đơn hàng E2E (UML 14)")
                inserted_images.add("14-seq")

            elif current_chapter == 2 and "UML 19: Vận đơn GHN tự động" in item_content and "19-seq" not in inserted_images:
                fig_count[2] += 1
                add_image_scaled(doc, os.path.join(UML_IMAGES_DIR, "19-sequence-ghn-shipment.png"),
                                 f"Hình 2.{fig_count[2]}: Sơ đồ tuần tự - Vận đơn GHN tự động & Webhook đồng bộ (UML 19)")
                fig_count[2] += 1
                add_image_scaled(doc, os.path.join(UML_IMAGES_DIR, "10-state-machine-shipment.png"),
                                 f"Hình 2.{fig_count[2]}: Sơ đồ máy trạng thái - Vòng đời Vận đơn Giao Hàng Nhanh (UML 10)")
                inserted_images.add("19-seq")

            elif current_chapter == 2 and "UML 20: Đăng ký & Xác thực OTP" in item_content and "20-seq" not in inserted_images:
                fig_count[2] += 1
                add_image_scaled(doc, os.path.join(UML_IMAGES_DIR, "20-sequence-auth-otp.png"),
                                 f"Hình 2.{fig_count[2]}: Sơ đồ tuần tự - Đăng ký & Xác thực OTP có giới hạn nghiêm ngặt (UML 20)")
                fig_count[2] += 1
                add_image_scaled(doc, os.path.join(UML_IMAGES_DIR, "21-flowchart-pusher-events.png"),
                                 f"Hình 2.{fig_count[2]}: Sơ đồ luồng sự kiện Realtime Pusher toàn hệ thống (UML 21)")
                inserted_images.add("20-seq")

            continue

        # Xử lý Đoạn văn bản thường (Body paragraph)
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.35
            p.paragraph_format.first_line_indent = Cm(1.0)

            parse_inline_formatting(p, line_strip, base_font_size=13, base_color=COLOR_BODY)
            i += 1


def generate_document(docx_path, toc_entries=None):
    """Tạo toàn bộ tệp DOCX hoàn chỉnh"""
    doc = docx.Document()

    # SECTION 1: Các trang đầu (Bìa, Lời cảm ơn, Nhận xét GVHD, Mục lục)
    s1 = doc.sections[0]
    s1.page_width = Cm(21.0)
    s1.page_height = Cm(29.7)
    s1.top_margin = Cm(2.0)
    s1.bottom_margin = Cm(2.0)
    s1.left_margin = Cm(3.0)
    s1.right_margin = Cm(2.0)

    build_outer_cover(doc)
    build_inner_cover(doc)
    build_acknowledgements(doc)
    build_supervisor_review(doc)
    build_table_of_contents(doc, toc_entries=toc_entries)

    # SECTION 2: Nội dung 5 chương chính
    s2 = doc.add_section(WD_SECTION.NEW_PAGE)
    s2.page_width = Cm(21.0)
    s2.page_height = Cm(29.7)
    s2.top_margin = Cm(2.0)
    s2.bottom_margin = Cm(2.0)
    s2.left_margin = Cm(3.0)
    s2.right_margin = Cm(2.0)

    # Ngắt liên kết Header & Footer với Section 1
    s2.header.is_linked_to_previous = False
    s2.footer.is_linked_to_previous = False

    # Header trang nội dung
    hp = s2.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run("Báo cáo Đồ án cơ sở 2 — Hệ thống Shop QR Payment")
    hrun.font.name = "Times New Roman"
    hrun.font.size = Pt(8.5)
    hrun.font.italic = True
    hrun.font.color.rgb = COLOR_GRAY

    # Footer trang nội dung (Đánh số trang 1, 2, 3...)
    fp = s2.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    frun = fp.add_run()
    frun.font.name = "Times New Roman"
    frun.font.size = Pt(10)
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    fp._p.append(fldSimple)

    # Thiết lập số trang bắt đầu từ 1
    sectPr = s2._sectPr
    pgNumType = OxmlElement('w:pgNumType')
    pgNumType.set(qn('w:start'), '1')
    sectPr.append(pgNumType)

    # Xây dựng nội dung 5 chương
    build_report_body(doc)

    doc.save(docx_path)
    print(f"  ✓ Đã ghi file DOCX thành công: {docx_path} ({os.path.getsize(docx_path):,} bytes)")


def convert_to_pdf(docx_path, out_dir):
    """Chuyển đổi DOCX sang PDF bằng LibreOffice headless"""
    cmd = [
        'libreoffice',
        '--headless',
        '--convert-to', 'pdf',
        docx_path,
        '--outdir', out_dir
    ]
    res = subprocess.run(cmd, capture_output=True, text=True)
    if res.returncode != 0:
        print(f"  [ERR] LibreOffice convert error: {res.stderr}")
        raise RuntimeError(f"LibreOffice failed: {res.stderr}")
    
    base_name = os.path.splitext(os.path.basename(docx_path))[0]
    pdf_path = os.path.join(out_dir, f"{base_name}.pdf")
    if os.path.exists(pdf_path):
        print(f"  ✓ Đã tạo PDF thành công: {pdf_path} ({os.path.getsize(pdf_path):,} bytes)")
        return pdf_path
    else:
        raise FileNotFoundError(f"PDF not found at {pdf_path}")


def extract_real_toc_pages(pdf_path):
    """Dùng PyMuPDF quét tìm vị trí chính xác của từng tiêu đề mục trong PDF"""
    import pymupdf
    pdf = pymupdf.open(pdf_path)
    total_pages = len(pdf)
    print(f"  ℹ Tổng số trang PDF hiện tại: {total_pages}")

    # Danh sách các tiêu đề cần dò
    titles_to_find = [
        ("CHƯƠNG I: TỔNG QUAN ĐỀ TÀI", 1),
        ("1.1. Giới thiệu", 2),
        ("1.1.1. Mục đích", 3),
        ("1.1.2. Quy ước tài liệu", 3),
        ("1.1.3. Đối tượng và gợi ý", 3),
        ("1.1.4. Phạm vi dự án", 3),
        ("1.2. Mô tả tổng quan", 2),
        ("1.2.1. Quan điểm về hệ thống", 3),
        ("1.2.2. Các lớp người dùng và đặc điểm", 3),
        ("1.2.3. Môi trường hoạt động", 3),
        ("1.2.4. Hạn chế về thiết kế và triển khai", 3),
        ("1.2.5. Tài liệu hướng dẫn người dùng", 3),
        ("1.2.6. Giả định và phụ thuộc", 3),
        ("1.3. Observation (Biên bản họp nhóm)", 2),
        ("CHƯƠNG II: MÔ HÌNH VÀ CHỨC NĂNG", 1),
        ("2.1. Các yêu cầu chức năng", 2),
        ("2.1.1. Các tác nhân", 3),
        ("2.1.2. Các tính năng của hệ thống", 3),
        ("2.1.3. Thành lập các chức năng", 3),
        ("2.1.4. Đặc tả use case", 3),
        ("2.2. Biểu đồ Use Case", 2),
        ("2.3. Class Diagram", 2),
        ("2.4. Activity Diagram", 2),
        ("2.5. Sequence Diagram", 2),
        ("2.6. Communication Diagram", 2),
        ("2.7. Mô hình cơ sở dữ liệu quan hệ", 2),
        ("CHƯƠNG III: XÂY DỰNG ỨNG DỤNG", 1),
        ("3.1. Kiến trúc hệ thống và công nghệ sử dụng", 2),
        ("3.2. Giao diện trang Khách hàng (Storefront)", 2),
        ("3.3. Giao diện trang Quản trị (Admin Bento Grid)", 2),
        ("3.4. Giao diện trang Nhân viên (Staff Portal)", 2),
        ("CHƯƠNG IV: KẾT LUẬN", 1),
        ("4.1. Ưu điểm của hệ thống", 2),
        ("4.2. Hạn chế của hệ thống", 2),
        ("4.3. Hướng phát triển trong tương lai", 2),
        ("4.4. Lời kết", 2),
        ("CHƯƠNG V: TÀI LIỆU THAM KHẢO", 1),
    ]

    # Tìm trang bắt đầu của Chương 1 trong PDF (để tính offset đánh số trang nội dung)
    ch1_page_idx = None
    for p_no in range(min(15, total_pages)):
        text = pdf[p_no].get_text()
        if "CHƯƠNG I: TỔNG QUAN ĐỀ TÀI" in text:
            ch1_page_idx = p_no
            break
    
    if ch1_page_idx is None:
        ch1_page_idx = 4  # Mặc định trang thứ 5
    
    print(f"  ℹ Chương 1 bắt đầu từ trang vật lý số {ch1_page_idx + 1} (Số trang nội dung: 1)")

    toc_results = []
    for title, level in titles_to_find:
        found_page_num = 1
        # Tìm kiếm từ trang Chương 1 trở đi
        for p_no in range(ch1_page_idx, total_pages):
            p_text = pdf[p_no].get_text()
            if title in p_text:
                # Trang hiển thị trong footer = (p_no - ch1_page_idx + 1)
                found_page_num = (p_no - ch1_page_idx + 1)
                break
        toc_results.append((title, level, found_page_num))

    return toc_results


def main():
    print("======================================================================")
    print("🚀 BẮT ĐẦU QUY TRÌNH XUẤT BẢN BÁO CÁO DACS2 CHUẨN VKU (.DOCX & .PDF)")
    print("======================================================================")

    # Đảm bảo thư mục đích tồn tại
    os.makedirs(ACADEMIC_DIR, exist_ok=True)

    # BƯỚC 1: Pass 1 - Tạo file DOCX với Mục lục sơ bộ
    print("\n--- BƯỚC 1: Tạo bản DOCX lần 1 (Pass 1) ---")
    generate_document(DOCX_OUT_PATH)

    # BƯỚC 2: Pass 1 - Biên dịch thử sang PDF để tính toán chính xác số trang
    print("\n--- BƯỚC 2: Biên dịch PDF lần 1 để định vị số trang chuẩn ---")
    pdf_temp = convert_to_pdf(DOCX_OUT_PATH, ACADEMIC_DIR)

    # BƯỚC 3: Dò tìm vị trí thực tế của tất cả các tiêu đề
    print("\n--- BƯỚC 3: Quét vị trí trang thực tế của từng tiêu đề mục ---")
    real_toc = extract_real_toc_pages(pdf_temp)
    print(f"  ✓ Đã trích xuất chính xác {len(real_toc)} mục lục!")
    for item in real_toc[:8]:
        print(f"    - {item[0]} -> Trang {item[2]}")

    # BƯỚC 4: Pass 2 - Tái xuất bản DOCX với Mục lục chính xác 100%
    print("\n--- BƯỚC 4: Tái xuất bản file DOCX chính thức (Pass 2 - Exact TOC) ---")
    generate_document(DOCX_OUT_PATH, toc_entries=real_toc)

    # BƯỚC 5: Pass 2 - Biên dịch PDF chính thức
    print("\n--- BƯỚC 5: Biên dịch PDF chính thức hoàn chỉnh ---")
    final_pdf = convert_to_pdf(DOCX_OUT_PATH, ACADEMIC_DIR)

    # BƯỚC 6: Sao chép đồng bộ vào Workspace và các thư mục phân phối
    print("\n--- BƯỚC 6: Đồng bộ hóa tài liệu thành phẩm vào hệ thống ---")
    dist_dirs = [
        "/home/nhonhoa/.hermes/kanban/boards/shop-qr-payment/workspaces/t_e9a3abdb",
        "/home/nhonhoa/Documents",
        "/home/nhonhoa/Downloads"
    ]
    for d in dist_dirs:
        os.makedirs(d, exist_ok=True)
        shutil.copy2(DOCX_OUT_PATH, os.path.join(d, "BAO_CAO_DACS2_SHOP_QR_PAYMENT.docx"))
        shutil.copy2(final_pdf, os.path.join(d, "BAO_CAO_DACS2_SHOP_QR_PAYMENT.pdf"))
        print(f"  ✓ Đã đồng bộ sang: {d}")

    # Kiểm tra kích thước và số trang thành phẩm
    import pymupdf
    pdf = pymupdf.open(final_pdf)
    print("\n======================================================================")
    print("🎉 HOÀN THÀNH XUẤT BẢN THÀNH CÔNG TÀI LIỆU BÁO CÁO DACS2!")
    print(f"  📄 File DOCX: {DOCX_OUT_PATH} ({os.path.getsize(DOCX_OUT_PATH):,} bytes)")
    print(f"  📕 File PDF:  {final_pdf} ({os.path.getsize(final_pdf):,} bytes, {len(pdf)} trang)")
    print("======================================================================")


if __name__ == "__main__":
    main()
