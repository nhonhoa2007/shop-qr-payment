#!/usr/bin/env python3
"""
Script xuất bản Báo cáo Đồ án cơ sở 2 (DACS2) cho dự án Shop QR Payment.
Tuân thủ nghiêm ngặt tiêu chuẩn học thuật VKU:
- Khổ giấy A4, Lề trang: Top 2.0cm, Bottom 2.0cm, Left 3.0cm (gáy), Right 2.0cm.
- Font Times New Roman, Size 13pt, Giãn dòng 1.3, Space after 6pt.
- Heading 1 (16pt Bold #1E3A8A, ngắt trang), Heading 2 (14pt Bold #1E3A8A), Heading 3 (13pt Bold #0F172A), Heading 4 (13pt Bold Italic #334155).
- Bảng biểu: Header #1E3A8A chữ trắng bold, viền #CBD5E1, xen kẽ dòng chẵn trắng / dòng lẻ #F8FAFC.
- Trang bìa chính (Outer Cover có viền đôi), Trang bìa phụ (Inner Cover), Lời cảm ơn, Nhận xét GVHD (kèm bảng điểm), Mục lục, Chương I -> V.
- Tích hợp 22 Use Cases và 19 bảng CSDL quan hệ.
- Giữ nguyên các placeholder cá nhân [...].
"""

import os
import re
import sys
import docx
from docx.shared import Cm, Mm, Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

DOCX_OUT = "/home/nhonhoa/shop-qr-payment/docs/academic/BAO_CAO_DACS2_SHOP_QR_PAYMENT.docx"
PDF_OUT_DIR = "/home/nhonhoa/shop-qr-payment/docs/academic"
MD_SRC = "/home/nhonhoa/shop-qr-payment/docs/academic/BAN_THAO_BAO_CAO_DACS2.md"

COLOR_PRIMARY = RGBColor(0x1E, 0x3A, 0x8A)      # Deep Navy #1E3A8A
COLOR_PRIMARY_HEX = "1E3A8A"
COLOR_SECONDARY = RGBColor(0x0F, 0x17, 0x2A)    # Dark Slate #0F172A
COLOR_MUTED = RGBColor(0x64, 0x74, 0x8B)        # Slate #64748B
COLOR_DARK = RGBColor(0x1E, 0x29, 0x3B)         # Slate 800
COLOR_BLACK = RGBColor(0x00, 0x00, 0x00)
COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_BORDER_HEX = "CBD5E1"                     # Slate 300
COLOR_ZEBRA_HEX = "F8FAFC"                      # Slate 50

FONT_NAME = "Times New Roman"

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    # Remove existing shd if any
    for child in tcPr.findall(qn('w:shd')):
        tcPr.remove(child)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_table_borders(table, color="CBD5E1", sz="4"):
    tblPr = table._tbl.tblPr
    for child in tblPr.findall(qn('w:tblBorders')):
        tblPr.remove(child)
    borders_xml = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders_xml)

def set_table_margins(table, top=120, bottom=120, left=160, right=160):
    tblPr = table._tbl.tblPr
    for child in tblPr.findall(qn('w:tblCellMar')):
        tblPr.remove(child)
    mar_xml = parse_xml(f'''
        <w:tblCellMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tblCellMar>
    ''')
    tblPr.append(mar_xml)

def set_cell_width(cell, width_cm):
    width_dxa = int(width_cm * 567)
    cell.width = Cm(width_cm)
    tcPr = cell._tc.get_or_add_tcPr()
    for child in tcPr.findall(qn('w:tcW')):
        tcPr.remove(child)
    tcW = parse_xml(f'<w:tcW {nsdecls("w")} w:w="{width_dxa}" w:type="dxa"/>')
    tcPr.append(tcW)

def parse_runs(text):
    """Phân rã chuỗi có định dạng markdown (**đậm**, *nghiêng*, `code`) thành danh sách runs."""
    tokens = []
    pattern = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)')
    parts = pattern.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            tokens.append((part[2:-2], {'bold': True}))
        elif part.startswith('*') and part.endswith('*'):
            tokens.append((part[1:-1], {'italic': True}))
        elif part.startswith('`') and part.endswith('`'):
            tokens.append((part[1:-1], {'code': True}))
        else:
            tokens.append((part, {}))
    return tokens

def add_formatted_runs_to_p(p, text, font_name=FONT_NAME, size_pt=13.0, default_color=None, default_bold=False, default_italic=False):
    tokens = parse_runs(text)
    for text_part, flags in tokens:
        run = p.add_run(text_part)
        run.font.name = font_name
        run.font.size = Pt(size_pt)
        is_bold = flags.get('bold', False) or default_bold
        is_italic = flags.get('italic', False) or default_italic
        if is_bold:
            run.bold = True
        if is_italic:
            run.italic = True
        if flags.get('code', False):
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        elif default_color:
            run.font.color.rgb = default_color
        else:
            run.font.color.rgb = COLOR_BLACK

def apply_base_page_setup(section):
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(3.0)   # Lề gáy sách 3.0cm
    section.right_margin = Cm(2.0)

def add_page_number_to_run(run):
    fldChar1 = parse_xml(rf'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    instrText = parse_xml(rf'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>')
    fldChar2 = parse_xml(rf'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>')
    fldChar3 = parse_xml(rf'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def build_cover_page(doc, is_outer=True):
    """Xây dựng trang bìa chính hoặc trang bìa phụ theo chuẩn VKU."""
    # Đoạn 1: Tên Trường
    p_school = doc.add_paragraph()
    p_school.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_school.paragraph_format.space_before = Pt(12)
    p_school.paragraph_format.space_after = Pt(2)
    run_school = p_school.add_run("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN & TRUYỀN THÔNG VIỆT - HÀN")
    run_school.font.name = FONT_NAME
    run_school.font.size = Pt(12.5)
    run_school.bold = True
    run_school.font.color.rgb = COLOR_BLACK

    # Đoạn 2: Tên Khoa
    p_dept = doc.add_paragraph()
    p_dept.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_dept.paragraph_format.space_before = Pt(2)
    p_dept.paragraph_format.space_after = Pt(24)
    run_dept = p_dept.add_run("KHOA KHOA HỌC MÁY TÍNH")
    run_dept.font.name = FONT_NAME
    run_dept.font.size = Pt(12.5)
    run_dept.bold = True
    run_dept.font.color.rgb = COLOR_BLACK

    # Đường phân cách trang trí
    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_before = Pt(0)
    p_div.paragraph_format.space_after = Pt(40)
    run_div = p_div.add_run("━━━━━━━━━━━━━━━━━━━")
    run_div.font.name = FONT_NAME
    run_div.font.size = Pt(10)
    run_div.font.color.rgb = COLOR_PRIMARY

    # Đoạn 3: Tên Học phần / Báo cáo
    p_type = doc.add_paragraph()
    p_type.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_type.paragraph_format.space_before = Pt(20)
    p_type.paragraph_format.space_after = Pt(14)
    run_type = p_type.add_run("BÁO CÁO ĐỒ ÁN CƠ SỞ 2")
    run_type.font.name = FONT_NAME
    run_type.font.size = Pt(22)
    run_type.bold = True
    run_type.font.color.rgb = COLOR_PRIMARY

    # Đoạn 4: Tiêu đề Đề tài
    p_label = doc.add_paragraph()
    p_label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_label.paragraph_format.space_before = Pt(10)
    p_label.paragraph_format.space_after = Pt(6)
    run_label = p_label.add_run("ĐỀ TÀI:")
    run_label.font.name = FONT_NAME
    run_label.font.size = Pt(14)
    run_label.bold = True
    run_label.font.color.rgb = COLOR_PRIMARY

    p_topic = doc.add_paragraph()
    p_topic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_topic.paragraph_format.space_before = Pt(4)
    p_topic.paragraph_format.space_after = Pt(45)
    run_topic = p_topic.add_run("WEBSITE THƯƠNG MẠI ĐIỆN TỬ TÍCH HỢP THANH TOÁN VIETQR TỰ ĐỘNG\nVÀ QUẢN TRỊ SAAS BENTO GRID\n(SHOP QR PAYMENT)")
    run_topic.font.name = FONT_NAME
    run_topic.font.size = Pt(14)
    run_topic.bold = True
    run_topic.font.color.rgb = COLOR_PRIMARY

    # Bảng thông tin sinh viên & giảng viên (Căn lề đẹp không viền)
    tbl_info = doc.add_table(rows=4, cols=2)
    tbl_info.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell_width(tbl_info.cell(0, 0), 5.2)
    set_cell_width(tbl_info.cell(0, 1), 9.8)
    set_cell_width(tbl_info.cell(1, 0), 5.2)
    set_cell_width(tbl_info.cell(1, 1), 9.8)
    set_cell_width(tbl_info.cell(2, 0), 5.2)
    set_cell_width(tbl_info.cell(2, 1), 9.8)
    set_cell_width(tbl_info.cell(3, 0), 5.2)
    set_cell_width(tbl_info.cell(3, 1), 9.8)

    info_rows = [
        ("Sinh viên thực hiện :", "[Họ và tên sinh viên 1 - MSSV 1]\n[Họ và tên sinh viên 2 - MSSV 2]"),
        ("Lớp :", "[Lớp sinh hoạt / học phần]"),
        ("Giảng viên hướng dẫn :", "[Học hàm, Học vị & Họ tên Giảng viên hướng dẫn]"),
        ("", "")
    ]
    for row_idx, (col0, col1) in enumerate(info_rows[:3]):
        cell0 = tbl_info.cell(row_idx, 0)
        p0 = cell0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(3)
        p0.paragraph_format.space_after = Pt(3)
        r0 = p0.add_run(col0)
        r0.font.name = FONT_NAME
        r0.font.size = Pt(12.5)
        r0.bold = True

        cell1 = tbl_info.cell(row_idx, 1)
        p1 = cell1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(3)
        p1.paragraph_format.space_after = Pt(3)
        r1 = p1.add_run(col1)
        r1.font.name = FONT_NAME
        r1.font.size = Pt(12.5)
        r1.bold = False

    # Đoạn 5: Nơi thực hiện và thời gian
    p_loc = doc.add_paragraph()
    p_loc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_loc.paragraph_format.space_before = Pt(70)
    p_loc.paragraph_format.space_after = Pt(0)
    run_loc = p_loc.add_run("Đà Nẵng, tháng 10 năm 2026")
    run_loc.font.name = FONT_NAME
    run_loc.font.size = Pt(12)
    run_loc.italic = True
    run_loc.font.color.rgb = COLOR_BLACK

def build_acknowledgments_page(doc):
    """Trang Lời cảm ơn."""
    p_h = doc.add_paragraph()
    p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_h.paragraph_format.space_before = Pt(18)
    p_h.paragraph_format.space_after = Pt(16)
    r_h = p_h.add_run("LỜI CẢM ƠN")
    r_h.font.name = FONT_NAME
    r_h.font.size = Pt(16)
    r_h.bold = True
    r_h.font.color.rgb = COLOR_PRIMARY

    p_note = doc.add_paragraph()
    p_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_note.paragraph_format.space_before = Pt(0)
    p_note.paragraph_format.space_after = Pt(14)
    r_note = p_note.add_run("[Phần thông tin cá nhân: Sinh viên tự hoàn thành theo thực tế]")
    r_note.font.name = FONT_NAME
    r_note.font.size = Pt(11)
    r_note.italic = True
    r_note.font.color.rgb = COLOR_MUTED

    p_body1 = doc.add_paragraph()
    p_body1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_body1.paragraph_format.first_line_indent = Cm(1.0)
    p_body1.paragraph_format.line_spacing = 1.3
    p_body1.paragraph_format.space_after = Pt(8)
    add_formatted_runs_to_p(p_body1, "Em xin bày tỏ lòng biết ơn sâu sắc nhất đến Quý Thầy/Cô trường Đại học Công nghệ Thông tin & Truyền thông Việt - Hàn (VKU), đặc biệt là các Thầy/Cô Khoa Khoa học Máy tính đã tận tình giảng dạy, truyền đạt những kiến thức quý báu về nền tảng khoa học máy tính, kỹ thuật phần mềm và kiến trúc hệ thống trong suốt thời gian học tập.")

    p_body2 = doc.add_paragraph()
    p_body2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_body2.paragraph_format.first_line_indent = Cm(1.0)
    p_body2.paragraph_format.line_spacing = 1.3
    p_body2.paragraph_format.space_after = Pt(8)
    add_formatted_runs_to_p(p_body2, "Đặc biệt, em xin gửi lời tri ân chân thành nhất tới `[Học hàm, Học vị & Họ tên Giảng viên hướng dẫn]`, người đã trực tiếp hướng dẫn, định hướng chuyên môn và tận tâm đồng hành cùng em trong suốt quá trình triển khai Đồ án cơ sở 2. Những ý kiến đóng góp, chỉ bảo nghiêm cẩn và tâm huyết của Thầy/Cô là kim chỉ nam quan trọng giúp em hoàn thiện đồ án này một cách khoa học, chuyên nghiệp và ứng dụng thực tiễn cao.")

    p_body3 = doc.add_paragraph()
    p_body3.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_body3.paragraph_format.first_line_indent = Cm(1.0)
    p_body3.paragraph_format.line_spacing = 1.3
    p_body3.paragraph_format.space_after = Pt(16)
    add_formatted_runs_to_p(p_body3, "Mặc dù đã nỗ lực hết mình để xây dựng hệ thống hoàn chỉnh và vận hành ổn định, song do giới hạn về mặt thời gian và kinh nghiệm thực tế, báo cáo khó tránh khỏi những thiếu sót nhất định. Em rất mong nhận được những nhận xét, góp ý quý báu từ Quý Thầy/Cô trong hội đồng chấm đồ án để sản phẩm được hoàn thiện hơn nữa.")

    p_end = doc.add_paragraph()
    p_end.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_end.paragraph_format.space_before = Pt(14)
    p_end.paragraph_format.space_after = Pt(0)
    r_end = p_end.add_run("Em xin chân thành cảm ơn!")
    r_end.font.name = FONT_NAME
    r_end.font.size = Pt(13)
    r_end.italic = True
    r_end.bold = True

def build_evaluation_page(doc):
    """Trang Nhận xét của Giảng viên hướng dẫn kèm Bảng điểm tiêu chí."""
    p_h = doc.add_paragraph()
    p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_h.paragraph_format.space_before = Pt(18)
    p_h.paragraph_format.space_after = Pt(16)
    r_h = p_h.add_run("NHẬN XÉT CỦA GIẢNG VIÊN HƯỚNG DẪN")
    r_h.font.name = FONT_NAME
    r_h.font.size = Pt(16)
    r_h.bold = True
    r_h.font.color.rgb = COLOR_PRIMARY

    criteria = [
        ("1. Ý thức thực hiện đồ án của sinh viên:", 2),
        ("2. Đánh giá nội dung và kết quả đạt được:", 2),
        ("3. Khả năng ứng dụng và tính sáng tạo:", 2),
    ]

    for title, dot_lines in criteria:
        p_t = doc.add_paragraph()
        p_t.paragraph_format.space_before = Pt(8)
        p_t.paragraph_format.space_after = Pt(2)
        r_t = p_t.add_run(title)
        r_t.font.name = FONT_NAME
        r_t.font.size = Pt(13)
        r_t.bold = True

        for _ in range(dot_lines):
            p_dot = doc.add_paragraph()
            p_dot.paragraph_format.space_before = Pt(2)
            p_dot.paragraph_format.space_after = Pt(2)
            r_dot = p_dot.add_run("…………………………………………………………………………………………………………")
            r_dot.font.name = FONT_NAME
            r_dot.font.size = Pt(11)
            r_dot.font.color.rgb = COLOR_MUTED

    # Mục 4: Bảng điểm đánh giá chi tiết
    p_t4 = doc.add_paragraph()
    p_t4.paragraph_format.space_before = Pt(8)
    p_t4.paragraph_format.space_after = Pt(4)
    r_t4 = p_t4.add_run("4. Điểm đánh giá chi tiết:")
    r_t4.font.name = FONT_NAME
    r_t4.font.size = Pt(13)
    r_t4.bold = True

    eval_data = [
        ("1", "Ý thức thực hiện, tinh thần học hỏi và thái độ", "2.0", ""),
        ("2", "Khả năng thu thập, nghiên cứu tài liệu và phân tích yêu cầu", "2.0", ""),
        ("3", "Mức độ hoàn thành khối lượng công việc và chất lượng phần mềm", "4.0", ""),
        ("4", "Bố cục, hình thức trình bày quyển báo cáo và thuyết minh", "2.0", ""),
        ("Tổng", "Tổng điểm đánh giá", "10.0", "")
    ]

    col_widths = [1.2, 9.8, 2.5, 2.5]
    table = doc.add_table(rows=len(eval_data)+1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, COLOR_BORDER_HEX, "4")
    set_table_margins(table, top=100, bottom=100, left=140, right=140)

    # Header
    headers = ["STT", "Tiêu chí đánh giá", "Điểm tối đa", "Điểm đánh giá"]
    hdr_row = table.rows[0]
    hdr_row._tr.get_or_add_trPr().append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    hdr_row._tr.get_or_add_trPr().append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    for col_idx, h_text in enumerate(headers):
        cell = hdr_row.cells[col_idx]
        set_cell_background(cell, COLOR_PRIMARY_HEX)
        set_cell_width(cell, col_widths[col_idx])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(h_text)
        r.font.name = FONT_NAME
        r.font.size = Pt(11)
        r.bold = True
        r.font.color.rgb = COLOR_WHITE

    # Body rows
    for r_idx, row_values in enumerate(eval_data, start=1):
        row = table.rows[r_idx]
        row._tr.get_or_add_trPr().append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        is_total = (r_idx == len(eval_data))
        bg_color = COLOR_ZEBRA_HEX if (r_idx % 2 == 1 and not is_total) else "FFFFFF"
        if is_total:
            bg_color = "F1F5F9"

        for col_idx, val in enumerate(row_values):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg_color)
            set_cell_width(cell, col_widths[col_idx])
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            if col_idx in (0, 2, 3):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(val)
            r.font.name = FONT_NAME
            r.font.size = Pt(11)
            if is_total or col_idx == 0:
                r.bold = True

    # Ký tên
    p_sign_date = doc.add_paragraph()
    p_sign_date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sign_date.paragraph_format.space_before = Pt(12)
    p_sign_date.paragraph_format.space_after = Pt(2)
    r_sd = p_sign_date.add_run("Đà Nẵng, ngày … tháng … năm 2026")
    r_sd.font.name = FONT_NAME
    r_sd.font.size = Pt(12)
    r_sd.italic = True

    p_sign_title = doc.add_paragraph()
    p_sign_title.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sign_title.paragraph_format.space_before = Pt(2)
    p_sign_title.paragraph_format.space_after = Pt(2)
    r_st = p_sign_title.add_run("Giảng viên hướng dẫn")
    r_st.font.name = FONT_NAME
    r_st.font.size = Pt(12.5)
    r_st.bold = True

    p_sign_sub = doc.add_paragraph()
    p_sign_sub.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sign_sub.paragraph_format.space_before = Pt(2)
    p_sign_sub.paragraph_format.space_after = Pt(10)
    r_ss = p_sign_sub.add_run("(Ký và ghi rõ họ tên)")
    r_ss.font.name = FONT_NAME
    r_ss.font.size = Pt(11)
    r_ss.italic = True
    r_ss.font.color.rgb = COLOR_MUTED

def build_toc_page(doc):
    """Trang Mục lục hiển thị đầy đủ cây cấu trúc 5 chương và 22 Use Cases."""
    p_h = doc.add_paragraph()
    p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_h.paragraph_format.space_before = Pt(18)
    p_h.paragraph_format.space_after = Pt(16)
    r_h = p_h.add_run("MỤC LỤC")
    r_h.font.name = FONT_NAME
    r_h.font.size = Pt(16)
    r_h.bold = True
    r_h.font.color.rgb = COLOR_PRIMARY

    toc_entries = [
        (1, "CHƯƠNG I: TỔNG QUAN ĐỀ TÀI"),
        (2, "1.1. Giới thiệu"),
        (3, "1.1.1. Mục đích"),
        (3, "1.1.2. Quy ước tài liệu"),
        (3, "1.1.3. Đối tượng và gợi ý"),
        (3, "1.1.4. Phạm vi dự án"),
        (2, "1.2. Mô tả tổng quan"),
        (3, "1.2.1. Quan điểm về hệ thống"),
        (3, "1.2.2. Các lớp người dùng và đặc điểm"),
        (3, "1.2.3. Môi trường hoạt động"),
        (3, "1.2.4. Hạn chế về thiết kế và triển khai"),
        (3, "1.2.5. Tài liệu hướng dẫn người dùng"),
        (3, "1.2.6. Giả định và phụ thuộc"),
        (2, "1.3. Observation (Biên bản họp nhóm)"),
        (1, "CHƯƠNG II: MÔ HÌNH VÀ CHỨC NĂNG"),
        (2, "2.1. Các yêu cầu chức năng"),
        (3, "2.1.1. Các tác nhân"),
        (3, "2.1.2. Các tính năng của hệ thống"),
        (3, "2.1.3. Thành lập các chức năng"),
        (3, "2.1.4. Đặc tả use case"),
        (4, "2.1.4.1. Đặc tả UC01: Đăng nhập hệ thống"),
        (4, "2.1.4.2. Đặc tả UC02: Đăng xuất"),
        (4, "2.1.4.3. Đặc tả UC03: Đăng ký tài khoản mới"),
        (4, "2.1.4.4. Đặc tả UC04: Xác thực mã OTP qua Email"),
        (4, "2.1.4.5. Đặc tả UC05: Khôi phục mật khẩu tài khoản"),
        (4, "2.1.4.6. Đặc tả UC06: Xem danh mục và Tìm kiếm sản phẩm"),
        (4, "2.1.4.7. Đặc tả UC07: Xem chi tiết sản phẩm và Lựa chọn biến thể SKU"),
        (4, "2.1.4.8. Đặc tả UC08: Quản lý giỏ hàng"),
        (4, "2.1.4.9. Đặc tả UC09: Áp dụng mã giảm giá Coupon"),
        (4, "2.1.4.10. Đặc tả UC10: Định vị GPS tự động điền địa chỉ giao hàng"),
        (4, "2.1.4.11. Đặc tả UC11: Đặt hàng và Thanh toán qua mã VietQR động"),
        (4, "2.1.4.12. Đặc tả UC12: Thanh toán đơn hàng bằng Ví Shop nội bộ"),
        (4, "2.1.4.13. Đặc tả UC13: Nạp tiền vào Ví Shop qua VietQR tĩnh"),
        (4, "2.1.4.14. Đặc tả UC14: Hủy đơn hàng và Hoàn tiền tự động vào Ví"),
        (4, "2.1.4.15. Đặc tả UC15: Chat tư vấn trực tuyến Realtime theo đơn hàng"),
        (4, "2.1.4.16. Đặc tả UC16: Đánh giá và Bình luận sản phẩm đã mua"),
        (4, "2.1.4.17. Đặc tả UC17: Quản lý danh sách Yêu thích (Wishlist)"),
        (4, "2.1.4.18. Đặc tả UC18: Quản lý đơn hàng và Duyệt trạng thái FSM"),
        (4, "2.1.4.19. Đặc tả UC19: Tạo vận đơn và Đồng bộ lộ trình Giao Hàng Nhanh (GHN)"),
        (4, "2.1.4.20. Đặc tả UC20: Quản lý Sản phẩm, Tồn kho và Ma trận Biến thể"),
        (4, "2.1.4.21. Đặc tả UC21: Đối soát thủ công giao dịch ngân hàng VietQR"),
        (4, "2.1.4.22. Đặc tả UC22: Phân quyền RBAC và Quản lý Người dùng"),
        (2, "2.2. Biểu đồ Use Case"),
        (3, "2.2.1. Biểu đồ Use Case tổng quát"),
        (3, "2.2.2. Biểu đồ Use Case dành cho Khách hàng"),
        (3, "2.2.3. Biểu đồ Use Case dành cho Nhân viên (Staff)"),
        (3, "2.2.4. Biểu đồ Use Case dành cho Quản trị viên (Admin)"),
        (2, "2.3. Class Diagram"),
        (2, "2.4. Activity Diagram"),
        (2, "2.5. Sequence Diagram"),
        (2, "2.6. Communication Diagram"),
        (2, "2.7. Mô hình cơ sở dữ liệu quan hệ"),
        (3, "2.7.1. Danh sách các bảng dữ liệu"),
        (3, "2.7.2. Bảng User (Người dùng)"),
        (3, "2.7.3. Bảng StaffPermission (Ma trận quyền nhân viên)"),
        (3, "2.7.4. Bảng OtpCode (Mã xác thực OTP)"),
        (3, "2.7.5. Bảng Product (Sản phẩm)"),
        (3, "2.7.6. Bảng ProductVariant (Biến thể sản phẩm)"),
        (3, "2.7.7. Bảng Order (Đơn hàng)"),
        (3, "2.7.8. Bảng OrderItem (Chi tiết mặt hàng trong đơn)"),
        (3, "2.7.9. Bảng Coupon (Mã giảm giá)"),
        (3, "2.7.10. Bảng CouponUsage (Lịch sử dùng mã giảm giá)"),
        (3, "2.7.11. Bảng Review (Đánh giá sản phẩm)"),
        (3, "2.7.12. Bảng Wishlist (Danh sách yêu thích)"),
        (3, "2.7.13. Bảng Transaction (Giao dịch ngân hàng VietQR)"),
        (3, "2.7.14. Bảng ChatRoom (Phòng chat đơn hàng)"),
        (3, "2.7.15. Bảng ChatRoomParticipant (Thành viên tham gia phòng chat)"),
        (3, "2.7.16. Bảng Message (Tin nhắn chat)"),
        (3, "2.7.17. Bảng Notification (Thông báo người dùng)"),
        (3, "2.7.18. Bảng Shipment (Vận đơn Giao Hàng Nhanh)"),
        (3, "2.7.19. Bảng UserWallet (Ví điện tử nội bộ)"),
        (3, "2.7.20. Bảng WalletTransaction (Lịch sử biến động số dư ví)"),
        (1, "CHƯƠNG III: XÂY DỰNG ỨNG DỤNG"),
        (2, "3.1. Kiến trúc hệ thống và công nghệ sử dụng"),
        (3, "3.1.1. Kiến trúc phân tầng sạch (Clean Architecture 3 lớp)"),
        (3, "3.1.2. Công nghệ và thư viện cốt lõi"),
        (2, "3.2. Giao diện trang Khách hàng (Storefront)"),
        (3, "3.2.1. Trang chủ và Danh mục sản phẩm (Catalog & Filtering)"),
        (3, "3.2.2. Trang Chi tiết sản phẩm & Lựa chọn biến thể SKU"),
        (3, "3.2.3. Trang Giỏ hàng trực tuyến"),
        (3, "3.2.4. Trang Đặt hàng (Checkout) & Định vị GPS"),
        (3, "3.2.5. Trang Thanh toán VietQR động"),
        (3, "3.2.6. Trang Lịch sử Đơn hàng & Tra cứu Vận đơn GHN"),
        (3, "3.2.7. Trang Quản lý Ví Shop & Nạp tiền VietQR"),
        (3, "3.2.8. Giao diện Chat trực tuyến Realtime"),
        (3, "3.2.9. Trang Đánh giá sản phẩm đã mua"),
        (3, "3.2.10. Trang Quản lý Danh sách Yêu thích (Wishlist)"),
        (3, "3.2.11. Trang Đăng ký, Xác thực OTP & Đăng nhập"),
        (2, "3.3. Giao diện trang Quản trị (Admin Bento Grid)"),
        (3, "3.3.1. Dashboard phân tích doanh thu Realtime (Bento Grid)"),
        (3, "3.3.2. Quản lý Đơn hàng & Duyệt FSM"),
        (3, "3.3.3. Quản lý Sản phẩm, Tồn kho & Ma trận Biến thể"),
        (3, "3.3.4. Quản lý Vận đơn Giao Hàng Nhanh"),
        (3, "3.3.5. Đối soát Giao dịch Ngân hàng VietQR Thủ công"),
        (3, "3.3.6. Quản lý Mã giảm giá Coupon"),
        (3, "3.3.7. Quản lý Người dùng & Ma trận Phân quyền STAFF"),
        (3, "3.3.8. Trung tâm Hỗ trợ Chat CSKH"),
        (3, "3.3.9. Kiểm duyệt & Phản hồi Đánh giá Khách hàng"),
        (2, "3.4. Giao diện trang Nhân viên (Staff Portal)"),
        (1, "CHƯƠNG IV: KẾT LUẬN"),
        (2, "4.1. Ưu điểm của hệ thống"),
        (2, "4.2. Hạn chế của hệ thống"),
        (2, "4.3. Hướng phát triển trong tương lai"),
        (2, "4.4. Lời kết"),
        (1, "CHƯƠNG V: TÀI LIỆU THAM KHẢO")
    ]

    for level, title in toc_entries:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.15
        if level == 1:
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Cm(0)
            r = p.add_run(title)
            r.font.name = FONT_NAME
            r.font.size = Pt(12)
            r.bold = True
            r.font.color.rgb = COLOR_PRIMARY
        elif level == 2:
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.left_indent = Cm(0.6)
            r = p.add_run(title)
            r.font.name = FONT_NAME
            r.font.size = Pt(11)
            r.bold = True
            r.font.color.rgb = COLOR_DARK
        elif level == 3:
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.left_indent = Cm(1.2)
            r = p.add_run(title)
            r.font.name = FONT_NAME
            r.font.size = Pt(10.5)
            r.font.color.rgb = COLOR_BLACK
        else: # level 4
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.left_indent = Cm(1.8)
            r = p.add_run(title)
            r.font.name = FONT_NAME
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

def render_table_block(doc, table_lines):
    """Xử lý và dựng bảng dữ liệu từ Markdown block."""
    raw_rows = []
    for line in table_lines:
        line_clean = line.strip()
        if not line_clean or not line_clean.startswith('|'):
            continue
        cells = [c.strip() for c in line_clean.strip('|').split('|')]
        # Bỏ qua dòng phân cách | :--- | :--- |
        if all(re.match(r'^:?-+:?$', c) for c in cells):
            continue
        raw_rows.append(cells)

    if not raw_rows:
        return

    num_cols = len(raw_rows[0])
    hdr_cells = raw_rows[0]

    # Quyết định chiều rộng từng cột dựa trên loại bảng (tổng chiều rộng trang chuẩn = 16.0cm)
    if num_cols == 2:
        if any("Thuộc tính" in c for c in hdr_cells):
            col_widths = [3.5, 12.5]
        elif any("Action" in c for c in hdr_cells):
            col_widths = [8.0, 8.0]
        else:
            col_widths = [4.5, 11.5]
    elif num_cols == 4:
        if any("Mã UC" in c for c in hdr_cells):
            col_widths = [1.8, 4.5, 4.0, 5.7]
        elif any("Tên bảng" in c or "Tên Model" in c for c in hdr_cells):
            col_widths = [1.2, 3.8, 3.8, 7.2]
        elif any("Tiêu chí" in c for c in hdr_cells):
            col_widths = [1.2, 9.8, 2.5, 2.5]
        else:
            col_widths = [2.0, 4.5, 4.5, 5.0]
    elif num_cols == 5:
        # Bảng từ điển dữ liệu CSDL 18-19 bảng
        col_widths = [1.0, 3.2, 2.8, 3.8, 5.2]
    else:
        # Phân bổ đều cho các bảng khác
        col_widths = [16.0 / num_cols] * num_cols

    table = doc.add_table(rows=len(raw_rows), cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, COLOR_BORDER_HEX, "4")
    set_table_margins(table, top=100, bottom=100, left=140, right=140)

    # Dòng Header
    hdr_row = table.rows[0]
    hdr_row._tr.get_or_add_trPr().append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    hdr_row._tr.get_or_add_trPr().append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    for c_idx in range(num_cols):
        cell = hdr_row.cells[c_idx]
        set_cell_background(cell, COLOR_PRIMARY_HEX)
        set_cell_width(cell, col_widths[c_idx])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        text_val = hdr_cells[c_idx] if c_idx < len(hdr_cells) else ""
        r = p.add_run(text_val)
        r.font.name = FONT_NAME
        r.font.size = Pt(11)
        r.bold = True
        r.font.color.rgb = COLOR_WHITE

    # Các dòng nội dung (Zebra striping)
    for r_idx in range(1, len(raw_rows)):
        row = table.rows[r_idx]
        row._tr.get_or_add_trPr().append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        # Dòng lẻ #F8FAFC, dòng chẵn trắng
        bg_color = COLOR_ZEBRA_HEX if (r_idx % 2 == 1) else "FFFFFF"
        row_vals = raw_rows[r_idx]

        for c_idx in range(num_cols):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg_color)
            set_cell_width(cell, col_widths[c_idx])
            cell_text = row_vals[c_idx] if c_idx < len(row_vals) else ""

            # Xử lý xuống dòng <br> trong ô
            sub_paragraphs = cell_text.split('<br>')
            for sub_idx, sub_p_text in enumerate(sub_paragraphs):
                if sub_idx == 0:
                    p = cell.paragraphs[0]
                else:
                    p = cell.add_paragraph()
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15

                # Căn lề: Cột STT hoặc Mã hoặc Boolean căn giữa, còn lại căn trái
                if num_cols in (4, 5) and c_idx in (0,):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif num_cols == 2 and any("Thuộc tính" in c for c in hdr_cells) and c_idx == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

                # Size 10.5pt cho bảng để tránh tràn dòng
                font_sz = 10.5 if num_cols in (2, 5) else 11
                add_formatted_runs_to_p(p, sub_p_text.strip(), font_name=FONT_NAME, size_pt=font_sz)

    # Thêm khoảng trống nhỏ sau bảng
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(4)

def add_figure_image(doc, img_path, caption_text, width_cm=15.5):
    """Chèn hình ảnh sơ đồ/mockup và chú thích hình ảnh chuẩn học thuật."""
    if not os.path.exists(img_path):
        return
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(6)
    p_img.paragraph_format.space_after = Pt(2)
    run_img = p_img.add_run()
    run_img.add_picture(img_path, width=Cm(width_cm))

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(10)
    r_cap = p_cap.add_run(caption_text)
    r_cap.font.name = FONT_NAME
    r_cap.font.size = Pt(11)
    r_cap.italic = True
    r_cap.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

def generate_document():
    print(f"[*] Khởi tạo Document và tải bản thảo từ {MD_SRC}...")
    with open(MD_SRC, "r", encoding="utf-8") as f:
        md_text = f.read()

    doc = docx.Document()

    # Cấu hình Style Normal
    style_normal = doc.styles['Normal']
    if hasattr(style_normal, 'font'):
        style_normal.font.name = FONT_NAME  # type: ignore
        style_normal.font.size = Pt(13)     # type: ignore
        style_normal.font.color.rgb = COLOR_BLACK  # type: ignore
    rPr = style_normal.element.get_or_add_rPr()
    rFonts = parse_xml(rf'<w:rFonts {nsdecls("w")} w:ascii="{FONT_NAME}" w:hAnsi="{FONT_NAME}" w:cs="{FONT_NAME}"/>')
    rPr.append(rFonts)

    # -------------------------------------------------------------
    # PHẦN 1: TRANG BÌA CHÍNH (OUTER COVER)
    # -------------------------------------------------------------
    print("[*] Tạo Trang bìa chính (Outer Cover)...")
    sec_cover = doc.sections[0]
    apply_base_page_setup(sec_cover)

    # Thêm viền đôi trang trí cho trang bìa chính
    pgBorders_xml = parse_xml(f'''
        <w:pgBorders {nsdecls("w")} w:offsetFrom="page">
            <w:top w:val="double" w:sz="12" w:space="24" w:color="{COLOR_PRIMARY_HEX}"/>
            <w:left w:val="double" w:sz="12" w:space="24" w:color="{COLOR_PRIMARY_HEX}"/>
            <w:bottom w:val="double" w:sz="12" w:space="24" w:color="{COLOR_PRIMARY_HEX}"/>
            <w:right w:val="double" w:sz="12" w:space="24" w:color="{COLOR_PRIMARY_HEX}"/>
        </w:pgBorders>
    ''')
    sec_cover._sectPr.append(pgBorders_xml)
    build_cover_page(doc, is_outer=True)

    # -------------------------------------------------------------
    # PHẦN 2: TRANG BÌA PHỤ (INNER COVER)
    # -------------------------------------------------------------
    print("[*] Tạo Trang bìa phụ (Inner Cover)...")
    sec_inner = doc.add_section(WD_SECTION.NEW_PAGE)
    apply_base_page_setup(sec_inner)
    # Gỡ bỏ viền trang từ section trước
    for b in sec_inner._sectPr.xpath('./w:pgBorders'):
        sec_inner._sectPr.remove(b)
    sec_inner.header.is_linked_to_previous = False
    sec_inner.footer.is_linked_to_previous = False
    build_cover_page(doc, is_outer=False)

    # -------------------------------------------------------------
    # PHẦN 3: LỜI CẢM ƠN
    # -------------------------------------------------------------
    print("[*] Tạo Trang Lời cảm ơn...")
    sec_front = doc.add_section(WD_SECTION.NEW_PAGE)
    apply_base_page_setup(sec_front)
    for b in sec_front._sectPr.xpath('./w:pgBorders'):
        sec_front._sectPr.remove(b)
    sec_front.header.is_linked_to_previous = False
    sec_front.footer.is_linked_to_previous = False
    build_acknowledgments_page(doc)

    # -------------------------------------------------------------
    # PHẦN 4: NHẬN XÉT CỦA GIẢNG VIÊN HƯỚNG DẪN
    # -------------------------------------------------------------
    print("[*] Tạo Trang Nhận xét của Giảng viên hướng dẫn...")
    doc.add_page_break()
    build_evaluation_page(doc)

    # -------------------------------------------------------------
    # PHẦN 5: MỤC LỤC
    # -------------------------------------------------------------
    print("[*] Tạo Trang Mục lục...")
    doc.add_page_break()
    build_toc_page(doc)

    # -------------------------------------------------------------
    # PHẦN 6: THÂN BÁO CÁO (CHƯƠNG I -> CHƯƠNG V)
    # -------------------------------------------------------------
    print("[*] Tạo Thân bài báo cáo (Chương I - V)...")
    sec_body = doc.add_section(WD_SECTION.NEW_PAGE)
    apply_base_page_setup(sec_body)
    for b in sec_body._sectPr.xpath('./w:pgBorders'):
        sec_body._sectPr.remove(b)
    sec_body.header.is_linked_to_previous = False
    sec_body.footer.is_linked_to_previous = False

    # Header đầu trang: Căn phải, 8.5pt, In nghiêng, Slate
    hdr_p = sec_body.header.paragraphs[0]
    hdr_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hdr_r = hdr_p.add_run("Báo cáo Đồ án cơ sở 2 — Hệ thống Shop QR Payment")
    hdr_r.font.name = FONT_NAME
    hdr_r.font.size = Pt(8.5)
    hdr_r.italic = True
    hdr_r.font.color.rgb = COLOR_MUTED

    # Footer chân trang: Căn giữa, Đánh số trang tự động bắt đầu từ 1
    ftr_p = sec_body.footer.paragraphs[0]
    ftr_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ftr_r = ftr_p.add_run()
    ftr_r.font.name = FONT_NAME
    ftr_r.font.size = Pt(10)
    ftr_r.font.color.rgb = COLOR_MUTED
    add_page_number_to_run(ftr_r)

    pgNumType = parse_xml(rf'<w:pgNumType {nsdecls("w")} w:start="1"/>')
    sec_body._sectPr.append(pgNumType)

    # Tách lấy phần nội dung bắt đầu từ Chương I
    ch1_pos = md_text.find("# CHƯƠNG I: TỔNG QUAN ĐỀ TÀI")
    if ch1_pos == -1:
        raise ValueError("Không tìm thấy tiêu đề '# CHƯƠNG I: TỔNG QUAN ĐỀ TÀI' trong markdown.")
    body_md = md_text[ch1_pos:]

    # Phân tích các khối markdown: Heading, Table, List, Paragraph
    lines = body_md.split('\n')
    i = 0
    total_lines = len(lines)
    is_first_ch = True

    while i < total_lines:
        line = lines[i]
        line_s = line.strip()

        # Dòng trống hoặc phân cách ---
        if not line_s or line_s == '---':
            i += 1
            continue

        # 1. BẢNG DỮ LIỆU (Bắt đầu bằng '|')
        if line_s.startswith('|'):
            table_lines = []
            while i < total_lines and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            render_table_block(doc, table_lines)
            continue

        # 2. HEADING 1 (Tên Chương)
        if line_s.startswith('# '):
            title = line_s[2:].strip().replace('`', '')
            if not is_first_ch:
                doc.add_page_break()
            is_first_ch = False

            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(title)
            r.font.name = FONT_NAME
            r.font.size = Pt(16)
            r.bold = True
            r.font.color.rgb = COLOR_PRIMARY
            i += 1
            continue

        # 3. HEADING 2 (Mục lớn: 1.1, 2.1, 3.1...)
        if line_s.startswith('## '):
            title = line_s[3:].strip().replace('`', '')
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(title)
            r.font.name = FONT_NAME
            r.font.size = Pt(14)
            r.bold = True
            r.font.color.rgb = COLOR_PRIMARY
            i += 1
            continue

        # 4. HEADING 3 (Mục con: 1.1.1, 2.1.1...)
        if line_s.startswith('### '):
            title = line_s[4:].strip().replace('`', '')
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(title)
            r.font.name = FONT_NAME
            r.font.size = Pt(13)
            r.bold = True
            r.font.color.rgb = COLOR_SECONDARY
            i += 1

            # Tự động chèn hình minh họa sơ đồ cho các mục tương ứng
            if "2.2.1. Biểu đồ Use Case tổng quát" in title:
                p_desc = doc.add_paragraph()
                p_desc.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p_desc.paragraph_format.first_line_indent = Cm(1.0)
                p_desc.paragraph_format.space_after = Pt(6)
                p_desc.paragraph_format.line_spacing = 1.3
                add_formatted_runs_to_p(p_desc, "Biểu đồ Use Case tổng quát thể hiện toàn cảnh ranh giới hệ thống phân chia giữa 2 phân hệ: Phân hệ Mua sắm (Storefront) và Phân hệ Quản trị SaaS & Tự động hóa, kết nối với 3 tác nhân người dùng và 3 tác nhân hệ thống bên ngoài.")
                add_figure_image(doc, "/home/nhonhoa/shop-qr-payment/docs/uml/images/01-use-case.png", "Hình 2.1: Sơ đồ Use Case tổng quát hệ thống Shop QR Payment")
                # Bỏ qua các dòng mô tả cũ nếu trùng
                if i < total_lines and "Biểu đồ Use Case tổng quát" in lines[i]:
                    i += 1
                if i < total_lines and "*(Tham chiếu chi tiết:" in lines[i]:
                    i += 1
            elif "3.3.1. Dashboard phân tích doanh thu Realtime" in title:
                # Chèn hình mockup Bento Grid
                add_figure_image(doc, "/home/nhonhoa/shop-qr-payment/public/admin-mockups/option-1-saas-bento.png", "Hình 3.1: Giao diện Bảng điều khiển Quản trị SaaS Bento Grid 21 chỉ số analytics")
            elif "3.3.2. Quản lý Đơn hàng & Duyệt FSM" in title:
                # Chèn hình mockup Operations Dual Panel
                add_figure_image(doc, "/home/nhonhoa/shop-qr-payment/public/admin-mockups/option-3-operations-dual-panel.png", "Hình 3.2: Giao diện Vận hành Đơn hàng và Vận chuyển GHN")
            continue

        # 5. HEADING 4 (Mục cấp 4: 2.1.4.1 Đặc tả UC01...)
        if line_s.startswith('#### '):
            title = line_s[5:].strip().replace('`', '')
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(title)
            r.font.name = FONT_NAME
            r.font.size = Pt(13)
            r.bold = True
            r.italic = True
            r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
            i += 1
            continue

        # 6. DANH SÁCH KHÔNG THỨ TỰ (Bullet list '- ')
        if line_s.startswith('- '):
            text_item = line_s[2:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(-0.4)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.3
            r_bullet = p.add_run("• ")
            r_bullet.font.name = FONT_NAME
            r_bullet.font.size = Pt(11)
            r_bullet.bold = True
            r_bullet.font.color.rgb = COLOR_PRIMARY
            add_formatted_runs_to_p(p, text_item)
            i += 1
            continue

        # 7. DANH SÁCH CÓ THỨ TỰ (Numbered list '1. ')
        num_match = re.match(r'^(\d+)\.\s+(.*)$', line_s)
        if num_match:
            num_str = num_match.group(1)
            text_item = num_match.group(2)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(-0.4)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.3
            r_num = p.add_run(f"{num_str}. ")
            r_num.font.name = FONT_NAME
            r_num.font.size = Pt(13)
            r_num.bold = True
            r_num.font.color.rgb = COLOR_SECONDARY
            add_formatted_runs_to_p(p, text_item)
            i += 1
            continue

        # 8. CÁC ĐOẠN ĐẶC BIỆT (Ghi chú, Tham chiếu sơ đồ, Ràng buộc)
        if line_s.startswith('*Ràng buộc độc nhất:*') or line_s.startswith('*(Tham chiếu chi tiết:'):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.2
            add_formatted_runs_to_p(p, line_s, size_pt=11.5, default_italic=True, default_color=RGBColor(0x33, 0x41, 0x55))
            i += 1
            continue

        # 9. ĐOẠN VĂN BẢN THƯỜNG (Body paragraph)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(1.0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.3
        add_formatted_runs_to_p(p, line_s)
        i += 1

    print(f"[*] Lưu file Word vào {DOCX_OUT}...")
    doc.save(DOCX_OUT)
    print(f"[+] Hoàn thành lưu DOCX: {DOCX_OUT} ({os.path.getsize(DOCX_OUT)} bytes)")

if __name__ == "__main__":
    generate_document()
